import streamlit as st
from streamlit_option_menu import option_menu
import os
import re
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import base64
import json
from io import BytesIO
import PyPDF2
import docx
from datetime import datetime, timedelta
from utils import extract_text_from_pdf, extract_text_from_docx, call_openai_api, generate_report_pdf, get_download_link, create_gauge_chart, analyze_writing_quality_chunked, load_case_database, run_caseconnect_analysis, repair_broken_words
from assessment_criteria import ASSESSMENT_AREAS, ASSESSMENT_CRITERIA, calculate_weighted_score, get_score_color, get_grade_label, SECTOR_MAPPING, PROMPT_EXCLUSION_INSTRUCTIONS, KCM_COMPETENCIES
from case_generator import (
    CASE_TYPES,
    run_source_processing,
    draft_case_section,
    draft_caselet,
    run_compliance_passes,
    generate_teaching_note,
    build_case_docx,
    build_case_pdf,
    build_case_txt,
    build_teaching_note_docx,
    apply_compliance_fixes,
    map_kcm_competencies,
    GeneratorAPIError,
)
import base64
import urllib.request
import urllib.error
from urllib.parse import urlparse
from db_models import (
    initialize_session_state,
    register_user,
    authenticate_user,
    login_user,
    logout_user,
    save_assessment,
    get_user_assessments,
    get_assessment,
    save_generated_case,
    get_user_generated_cases,
    get_generated_case,
    delete_generated_case,
)

# Set page configuration
st.set_page_config(
    page_title="CBC-India AGK Case Study Suite",
    page_icon="📊",
    layout="wide"
)

# Load case database for CaseConnect
CASE_DB_PATH = "attached_assets/Case_details_-_Sheet1_1774840900142.csv"
@st.cache_data
def get_case_database():
    cases, structured_text, case_count = load_case_database(CASE_DB_PATH)
    link_lookup = {}
    for c in cases:
        if c.get("igot_link"):
            link_lookup[c["title"].strip().lower()] = c["igot_link"]
    return structured_text, case_count, link_lookup

CASE_DATABASE_TEXT, CASE_COUNT, CASE_LINK_LOOKUP = get_case_database()

def _resolve_igot_link(case_title):
    import unicodedata
    if not case_title:
        return ""
    def normalize(s):
        s = unicodedata.normalize('NFKD', s).strip().lower()
        s = re.sub(r'[\u2013\u2014\u2015\u2212]', '-', s)
        s = re.sub(r'[\u2018\u2019\u201c\u201d]', "'", s)
        s = re.sub(r'[^\w\s-]', '', s)
        s = re.sub(r'\s+', ' ', s).strip()
        return s

    norm_title = normalize(case_title)
    if case_title.strip().lower() in CASE_LINK_LOOKUP:
        return CASE_LINK_LOOKUP[case_title.strip().lower()]
    for db_title, link in CASE_LINK_LOOKUP.items():
        if normalize(db_title) == norm_title:
            return link
    for db_title, link in CASE_LINK_LOOKUP.items():
        norm_db = normalize(db_title)
        if len(norm_title) > 15 and len(norm_db) > 15:
            if norm_title in norm_db or norm_db in norm_title:
                return link
    return ""

# Custom CSS to improve the appearance
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Poppins:wght@300;400;500;600;700&display=swap');
    
    * {
        font-family: 'Poppins', sans-serif !important;
    }
    
    .main .block-container {
        padding-top: 2rem;
        padding-bottom: 2rem;
    }
    
    /* Global button styling to standardize sizes */
    .stButton > button, .stDownloadButton > button {
        width: 100% !important;
        height: 50px !important;
        font-weight: 600 !important;
        border-radius: 8px !important;
        margin-bottom: 10px !important;
    }
    
    /* Criteria Tab styling with wrapping and padding */
    .stTabs [data-baseweb="tab-list"] {
        gap: 10px;
        padding: 10px 0px;
    }
    
    .stTabs [data-baseweb="tab"] {
        height: auto !important;
        min-height: 50px !important;
        white-space: normal !important;
        word-wrap: break-word !important;
        max-width: 200px !important;
        padding: 10px 15px !important;
        line-height: 1.2 !important;
        text-align: center !important;
        background-color: #f0f2f6 !important;
        border-radius: 8px 8px 0px 0px !important;
        font-family: 'Poppins', sans-serif !important;
    }
    
    .stTabs [aria-selected="true"] {
        background-color: #074fa5 !important;
        color: white !important;
    }

    /* Result card header wrapping */
    .card-header, .criteria-heading, h4 {
        white-space: normal !important;
        word-wrap: break-word !important;
        line-height: 1.3 !important;
    }
    
    h1, h2, h3 {
        color: #1E3A8A;
        font-family: 'Poppins', sans-serif !important;
    }
    
    .stTextInput, .stButton, .stSelectbox {
        font-family: 'Poppins', sans-serif !important;
    }
    
    /* Make sure all text elements use Poppins */
    p, span, div, label, button, select, input, textarea {
        font-family: 'Poppins', sans-serif !important;
    }
    .card {
        border: 1px solid #D1D5DB;
        border-radius: 8px;
        padding: 16px;
        margin-bottom: 16px;
        background-color: #EFF6FF;
        box-shadow: 0 2px 4px rgba(0,0,0,0.12);
        color: #333333;
    }
    .criteria-heading {
        font-weight: bold;
        font-size: 1.1em;
        padding: 8px;
        margin-bottom: 12px;
        background-color: #DBEAFE;
        border-radius: 4px;
        color: #1E40AF;
    }
    .metrics-container {
        display: flex;
        justify-content: space-around;
        margin: 20px 0;
    }
    .metric-card {
        text-align: center;
        padding: 10px;
        border-radius: 5px;
        background-color: #F9FAFB;
        box-shadow: 0 1px 2px rgba(0,0,0,0.05);
    }
    /* Auth form styling */
    .auth-form {
        max-width: 450px;
        margin: 0 auto;
        padding: 20px;
        border-radius: 10px;
        box-shadow: 0 4px 6px rgba(0, 0, 0, 0.1);
        background-color: #F9FAFB;
    }
    .auth-form h3 {
        text-align: center;
        margin-bottom: 20px;
        color: #1E3A8A;
    }
    .auth-form .stButton>button {
        width: 100%;
        margin-top: 10px;
    }
    /* History card styling */
    .history-card {
        border: 1px solid #E5E7EB;
        border-radius: 8px;
        padding: 15px;
        margin-bottom: 15px;
        background-color: white;
        transition: transform 0.2s, box-shadow 0.2s;
    }
    .history-card:hover {
        transform: translateY(-2px);
        box-shadow: 0 4px 6px rgba(0, 0, 0, 0.1);
    }
    .history-card h4 {
        margin-top: 0;
        color: #1E3A8A;
    }
    .history-card p {
        margin-bottom: 5px;
        color: #4B5563;
    }
    .history-card .date {
        color: #6B7280;
        font-size: 0.9em;
    }
    /* User profile section */
    .user-profile {
        display: flex;
        align-items: center;
        margin-bottom: 20px;
        padding: 10px;
        background-color: #F0F9FF;
        border-radius: 8px;
    }
    .user-profile .avatar {
        width: 40px;
        height: 40px;
        border-radius: 50%;
        background-color: #3B82F6;
        color: white;
        display: flex;
        align-items: center;
        justify-content: center;
        font-weight: bold;
        margin-right: 10px;
    }
    .user-profile .info {
        flex-grow: 1;
    }
    .user-profile .info p {
        margin: 0;
    }
    .user-profile .actions {
        margin-left: 10px;
    }
    /* Score badge styling */
    .score-badge {
        display: inline-block;
        padding: 4px 12px;
        border-radius: 12px;
        font-weight: bold;
        font-size: 0.9em;
    }
    .score-excellent {
        background-color: #D1FAE5;
        color: #065F46;
    }
    .score-good {
        background-color: #DBEAFE;
        color: #1E40AF;
    }
    .score-satisfactory {
        background-color: #FEF3C7;
        color: #92400E;
    }
    .score-needs-improvement {
        background-color: #FED7AA;
        color: #9A3412;
    }
    .score-poor {
        background-color: #FECACA;
        color: #991B1B;
    }
</style>
""", unsafe_allow_html=True)

# Initialize session state for authentication
initialize_session_state()

# Authentication and user management
def show_login_page():
    st.markdown('<div class="auth-form">', unsafe_allow_html=True)
    st.subheader("Login to Your Account")
    
    st.info("Please enter your credentials to access the Case Study Analyser")
    
    username = st.text_input("Username")
    password = st.text_input("Password", type="password")
    
    if st.button("Login", use_container_width=True):
        if username and password:
            success, user = authenticate_user(username, password)
            if success:
                login_user(user)
                st.success("Login successful!")
                st.rerun()
            else:
                st.error(user)  # Error message
        else:
            st.warning("Please enter both username and password")
    
    # Add note for authorized users
    st.markdown("""
    <div style="margin-top: 20px; padding: 10px; border-radius: 5px; background-color: #f8f9fa; font-size: 0.9em; color: #6c757d;">
    <strong>Note:</strong> This is a private application. Only authorized users can access this tool.
    </div>
    """, unsafe_allow_html=True)
    
    st.markdown('</div>', unsafe_allow_html=True)

def show_register_page():
    st.markdown('<div class="auth-form">', unsafe_allow_html=True)
    st.subheader("Create a New Account")
    
    username = st.text_input("Username")
    email = st.text_input("Email")
    password = st.text_input("Password", type="password")
    confirm_password = st.text_input("Confirm Password", type="password")
    
    col1, col2 = st.columns(2)
    with col1:
        if st.button("Register"):
            if not username or not email or not password:
                st.warning("Please fill in all fields")
            elif password != confirm_password:
                st.error("Passwords do not match")
            else:
                success, message = register_user(username, email, password)
                if success:
                    st.success(message)
                    st.session_state.show_register = False
                    st.rerun()
                else:
                    st.error(message)
    
    with col2:
        if st.button("Back to Login"):
            st.session_state.show_register = False
            st.rerun()
    
    st.markdown('</div>', unsafe_allow_html=True)

# Function to display header with logos
def display_header_with_logos():
    col1, col2, col3 = st.columns([1, 4, 1])
    with col1:
        if os.path.exists("attached_assets/cbc_logo_1770210514857.png"):
            st.image("attached_assets/cbc_logo_1770210514857.png", width=100)
    with col2:
        active = st.session_state.get("active_tool", None)
        if active == "analyser":
            st.title("Case Study Analyser")
        elif active == "caseconnect":
            st.title("CaseConnect")
        elif active == "generator":
            st.title("Case Study Generator")
        else:
            st.title("CBC-India AGK Case Study Suite")
    with col3:
        if os.path.exists("attached_assets/agk_logo_1770210514857.png"):
            st.image("attached_assets/agk_logo_1770210514857.png", width=100)

# Show login page or main app
if not st.session_state.logged_in:
    display_header_with_logos()
    show_login_page()
        
else:
    display_header_with_logos()
    
    active = st.session_state.get("active_tool", None)
    col1, col2 = st.columns([3, 1])
    with col1:
        if active == "analyser":
            st.markdown("""
            This tool helps analyze case studies against the CBC-India AGK Case Study Review Rubric.
            Upload your case study document, and the AI will evaluate it based on four key assessment areas:
            1. Structure, Chronology & Logical Flow (30%)
            2. Language, Citations & Factual Accuracy (30%)
            3. Alignment with Teaching Note, Sector & Competencies (15%)
            4. Overall Effectiveness & Impact (25%)
            """)
        elif active == "caseconnect":
            st.markdown(f"""
            CaseConnect helps you discover relevant governance case studies from the AGK repository 
            for your teaching programme. Upload your course outline and answer a few questions — 
            the tool will recommend cases aligned with your course design, competencies, and sector.
            
            *Referring to {CASE_COUNT} case studies from the AGK repository.*
            """)
        elif active == "generator":
            st.markdown("""
            The Case Study Generator helps authors draft AGK-quality case studies from raw source
            material such as interview transcripts, reports, news articles, websites and notes.
            Three case types are supported — Lesson-Drawing, Decision-Forcing, and Caselet — and
            the tool enforces British English, past tense, neutral tone, the AGK template structure,
            and an optional teaching note.
            """)
    
    with col2:
        st.markdown(f"""
        <div class="user-profile">
            <div class="avatar">{st.session_state.username[0].upper()}</div>
            <div class="info">
                <p class="text-[14px]"><strong>{st.session_state.username}</strong></p>
            </div>
        </div>
        """, unsafe_allow_html=True)
        
        if st.button("Logout", key="logout_button"):
            logout_user()
            st.rerun()

# Initialize session state variables if they don't exist
if 'active_tool' not in st.session_state:
    st.session_state.active_tool = None
if 'case_study_text' not in st.session_state:
    st.session_state.case_study_text = ""
if 'teaching_note_text' not in st.session_state:
    st.session_state.teaching_note_text = ""
if 'case_study_analysis' not in st.session_state:
    st.session_state.case_study_analysis = None
if 'case_study_summary' not in st.session_state:
    st.session_state.case_study_summary = ""
if 'assessment_results' not in st.session_state:
    st.session_state.assessment_results = {}
if 'document_name' not in st.session_state:
    st.session_state.document_name = ""
if 'teaching_note_name' not in st.session_state:
    st.session_state.teaching_note_name = ""
if 'loaded_from_history' not in st.session_state:
    st.session_state.loaded_from_history = False
if 'weighted_scores' not in st.session_state:
    st.session_state.weighted_scores = None
if 'competency_mapping' not in st.session_state:
    st.session_state.competency_mapping = None
if 'sector_tags' not in st.session_state:
    st.session_state.sector_tags = []
if 'sector_subthemes' not in st.session_state:
    st.session_state.sector_subthemes = {}
if 'keywords' not in st.session_state:
    st.session_state.keywords = []
if 'sidebar_tab' not in st.session_state:
    st.session_state.sidebar_tab = "New Assessment"
if 'caseconnect_results' not in st.session_state:
    st.session_state.caseconnect_results = None
if 'caseconnect_curriculum_text' not in st.session_state:
    st.session_state.caseconnect_curriculum_text = ""

# Case Study Generator session state
if 'gen_step' not in st.session_state:
    st.session_state.gen_step = 1
if 'gen_case_id' not in st.session_state:
    st.session_state.gen_case_id = None
if 'gen_case_type' not in st.session_state:
    st.session_state.gen_case_type = None
if 'gen_metadata' not in st.session_state:
    st.session_state.gen_metadata = {}
if 'gen_intent' not in st.session_state:
    st.session_state.gen_intent = {}
if 'gen_sources' not in st.session_state:
    st.session_state.gen_sources = []
if 'gen_processed' not in st.session_state:
    st.session_state.gen_processed = {}
if 'gen_sections' not in st.session_state:
    st.session_state.gen_sections = {}
if 'gen_compliance' not in st.session_state:
    st.session_state.gen_compliance = {}
if 'gen_teaching_note' not in st.session_state:
    st.session_state.gen_teaching_note = {}
if 'gen_kcm_mapping' not in st.session_state:
    st.session_state.gen_kcm_mapping = {}
if 'gen_final_documents' not in st.session_state:
    st.session_state.gen_final_documents = {}

# Sidebar with simple navigation (only for logged-in users)
with st.sidebar:
    if st.session_state.logged_in and st.session_state.get("active_tool") is not None:
        if st.button("← Back to Home", key="back_home_btn", use_container_width=True):
            st.session_state.active_tool = None
            st.session_state.caseconnect_results = None
            st.session_state.caseconnect_curriculum_text = ""
            st.rerun()

        if st.session_state.get("active_tool") == "generator":
            st.markdown("---")
            st.subheader("Case Generator")
            st.caption(f"Step {st.session_state.gen_step} of 8")
            ct = st.session_state.gen_case_type
            if ct:
                st.write(f"**Case type:** {CASE_TYPES[ct]['label']}")
            if st.session_state.gen_metadata.get("title"):
                st.write(f"**Title:** {st.session_state.gen_metadata['title']}")
            if st.button("Start a new case", key="gen_reset_btn", use_container_width=True):
                for k in [
                    "gen_step", "gen_case_id", "gen_case_type", "gen_metadata",
                    "gen_intent", "gen_sources", "gen_processed", "gen_sections",
                    "gen_compliance", "gen_teaching_note",
                    "gen_kcm_mapping", "gen_final_documents",
                ]:
                    if k in st.session_state:
                        del st.session_state[k]
                st.rerun()

            st.markdown("---")
            st.subheader("Saved drafts")
            try:
                saved_cases = get_user_generated_cases(st.session_state.user_id) if st.session_state.user_id else []
            except Exception:
                saved_cases = []
            if saved_cases:
                for sc in saved_cases[:10]:
                    label = (sc.title or "Untitled") + f"  ·  {sc.case_type or '?'}"
                    row = st.columns([5, 1])
                    with row[0]:
                        if st.button(label, key=f"load_gen_{sc.id}", use_container_width=True):
                            loaded = get_generated_case(sc.id, st.session_state.user_id)
                            if loaded:
                                st.session_state.gen_case_id = loaded["id"]
                                st.session_state.gen_case_type = loaded["case_type"]
                                st.session_state.gen_metadata = loaded["metadata"]
                                st.session_state.gen_intent = loaded["intent"]
                                st.session_state.gen_sources = loaded["sources"]
                                st.session_state.gen_processed = loaded["processed"]
                                st.session_state.gen_sections = loaded["sections"]
                                st.session_state.gen_compliance = loaded["compliance"]
                                st.session_state.gen_teaching_note = loaded["teaching_note"]
                                st.session_state.gen_kcm_mapping = loaded.get("kcm_mapping") or {}
                                st.session_state.gen_final_documents = loaded.get("final_documents") or {}
                                st.session_state.gen_step = 6 if loaded["sections"] else 5 if loaded["processed"] else 4
                                st.rerun()
                    with row[1]:
                        confirm_key = f"confirm_del_gen_{sc.id}"
                        if st.session_state.get(confirm_key):
                            if st.button("✓", key=f"do_del_gen_{sc.id}", help="Confirm delete", use_container_width=True):
                                try:
                                    if delete_generated_case(sc.id, st.session_state.user_id):
                                        if st.session_state.get("gen_case_id") == sc.id:
                                            st.session_state.gen_case_id = None
                                        st.session_state[confirm_key] = False
                                        st.rerun()
                                    else:
                                        st.warning("Could not delete (no permission or already removed).")
                                except Exception as e:
                                    st.error(f"Delete failed: {e}")
                        else:
                            if st.button("🗑", key=f"ask_del_gen_{sc.id}", help="Delete this draft", use_container_width=True):
                                st.session_state[confirm_key] = True
                                st.rerun()
            else:
                st.caption("No drafts saved yet.")

        st.markdown("---")

        if st.session_state.active_tool == "analyser":
            selected = option_menu(
                menu_title=None,
                options=["New Assessment", "History"],
                icons=["file-earmark-text", "clock-history"],
                menu_icon="cast",
                default_index=0 if st.session_state.sidebar_tab == "New Assessment" else 1,
                orientation="vertical",
                styles={
                    "container": {"padding": "5px", "background-color": "#fafafa"},
                    "icon": {"color": "#1E3A8A", "font-size": "18px"},
                    "nav-link": {
                        "font-size": "16px",
                        "text-align": "left",
                        "margin": "5px",
                        "padding": "10px 15px",
                        "--hover-color": "#eee",
                        "border-radius": "8px"
                    },
                    "nav-link-selected": {
                        "background-color": "#074fa5",
                        "color": "white",
                        "font-weight": "600"
                    },
                }
            )

            st.markdown("""
            <style>
            div[data-testid="stVerticalBlock"] div.nav-link-selected i {
                color: white !important;
            }
            </style>
            """, unsafe_allow_html=True)
            
            if selected != st.session_state.sidebar_tab:
                st.session_state.sidebar_tab = selected
                
                if selected == "New Assessment":
                    if not st.session_state.get('loaded_from_history', False):
                        st.session_state.case_study_text = ""
                        st.session_state.teaching_note_text = ""
                        st.session_state.case_study_analysis = None
                        st.session_state.case_study_summary = ""
                        st.session_state.assessment_results = {}
                        if 'recommendations' in st.session_state:
                            st.session_state.recommendations = ""
                        st.session_state.document_name = ""
                        st.session_state.teaching_note_name = ""
                        st.session_state.weighted_scores = None
                        st.session_state.competency_mapping = None
                        st.session_state.sector_tags = []
                        st.session_state.sector_subthemes = {}
                        st.session_state.keywords = []
                    else:
                        st.session_state.loaded_from_history = False
                elif selected == "History":
                    if 'selected_assessment' in st.session_state:
                        st.session_state.selected_assessment = None
                
                st.rerun()
            
            st.markdown("---")
        
        if st.session_state.sidebar_tab == "New Assessment" and st.session_state.get("active_tool") == "analyser":
            st.header("Upload Documents")
            
            st.subheader("1. Case Study Document")
            st.caption("Upload the case study document (PDF or DOCX)")
            case_study_file = st.file_uploader("Choose Case Study file", type=["pdf", "docx"], key="case_study_uploader")
            
            if case_study_file is not None:
                st.session_state.document_name = case_study_file.name
                
                # Process the uploaded file
                try:
                    if case_study_file.type == "application/pdf":
                        st.session_state.case_study_text = extract_text_from_pdf(case_study_file)
                    elif case_study_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                        st.session_state.case_study_text = extract_text_from_docx(case_study_file)
                    
                    st.success(f"Case Study: {case_study_file.name}")
                    text_length = len(st.session_state.case_study_text)
                    st.info(f"Extracted {text_length} characters")
                    
                except Exception as e:
                    st.error(f"Error processing case study: {str(e)}")
            
            st.markdown("---")
            
            st.subheader("2. Teaching Note Document")
            st.caption("Upload the corresponding Teaching Note (PDF or DOCX)")
            teaching_note_file = st.file_uploader("Choose Teaching Note file", type=["pdf", "docx"], key="teaching_note_uploader")
            
            if teaching_note_file is not None:
                st.session_state.teaching_note_name = teaching_note_file.name
                
                # Process the uploaded file
                try:
                    if teaching_note_file.type == "application/pdf":
                        st.session_state.teaching_note_text = extract_text_from_pdf(teaching_note_file)
                    elif teaching_note_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                        st.session_state.teaching_note_text = extract_text_from_docx(teaching_note_file)
                    
                    st.success(f"Teaching Note: {teaching_note_file.name}")
                    tn_length = len(st.session_state.teaching_note_text)
                    st.info(f"Extracted {tn_length} characters")
                    
                except Exception as e:
                    st.error(f"Error processing teaching note: {str(e)}")
            
            # Action buttons
            st.subheader("Analysis Actions")
            
            # Check if both documents are uploaded
            both_docs_uploaded = st.session_state.case_study_text and st.session_state.teaching_note_text
            
            if both_docs_uploaded:
                st.success("Both documents uploaded. Ready for assessment.")
                
                if st.button("Generate Summary"):
                    with st.spinner("Generating case study summary..."):
                        prompt = f"""
                        Please provide a concise summary of the following case study document. 
                        Focus on the main theme, key events, stakeholders involved, and outcomes:
                        
                        {st.session_state.case_study_text[:4000]}
                        
                        {PROMPT_EXCLUSION_INSTRUCTIONS}
                        """
                        
                        st.session_state.case_study_summary = call_openai_api(prompt)
                    
                    with st.spinner("Generating sector tags and keywords..."):
                        sector_prompt = f"""
                        You are an expert at classifying case studies by sector, sub-theme, and keywords.
                        
                        Based on the case study below, identify the applicable sectors, sub-themes, and keywords.
                        
                        === AVAILABLE SECTORS AND SUB-THEMES ===
                        {json.dumps(SECTOR_MAPPING, indent=2)}
                        
                        === CASE STUDY ===
                        {st.session_state.case_study_text[:5000]}
                        
                        {PROMPT_EXCLUSION_INSTRUCTIONS}
                        
                        Provide your analysis in the following JSON structure:
                        {{
                            "sectors": ["sector1", "sector2"],
                            "subthemes": {{"sector1": ["subtheme1", "subtheme2"], "sector2": ["subtheme1"]}},
                            "keywords": ["keyword1", "keyword2", "keyword3"]
                        }}
                        
                        Important:
                        - Select ONLY sectors from the provided SECTOR_MAPPING list
                        - For each selected sector, identify relevant sub-themes from that sector's sub-theme list
                        - Generate 5-10 keywords that a general public audience would use to search for this case study
                        - Keywords should be simple, commonly used terms
                        """
                        
                        tags_result = call_openai_api(sector_prompt, response_format="json_object", temperature=0.1, seed=42)
                        if isinstance(tags_result, dict):
                            st.session_state.sector_tags = tags_result.get("sectors", [])
                            st.session_state.sector_subthemes = tags_result.get("subthemes", {})
                            st.session_state.keywords = tags_result.get("keywords", [])
                
                if st.button("Perform Full Assessment"):
                    from concurrent.futures import ThreadPoolExecutor, as_completed
                    progress_text = st.empty()
                    progress_bar = st.progress(0)
                    
                    st.session_state.assessment_results = {}
                    st.session_state.writing_findings = []
                    
                    total_steps = len(ASSESSMENT_AREAS) + 3
                    completed_steps = 0

                    case_text = st.session_state.case_study_text
                    tn_text = st.session_state.teaching_note_text

                    def _build_criterion_prompt(area_id, criterion_info):
                        is_informational = criterion_info.get('informational', False)
                        max_score = criterion_info.get('max_score', 3) if not is_informational else 0
                        scoring_logic = criterion_info.get('scoring_logic', '')
                        agent_prompt = criterion_info.get('prompt', criterion_info['description'])
                        requires_tn = criterion_info.get('requires_teaching_note', False)
                        
                        if is_informational:
                            return f"""
                            You are a case study evaluation expert using the CBC-India AGK Case Study Review Rubric.
                            
                            Provide a detailed narrative analysis for the following criterion. Do NOT provide a score — this is an informational assessment only.
                            
                            **Criterion:** {criterion_info['name']}
                            **Description:** {criterion_info['description']}
                            **Analysis Task:** {agent_prompt}
                            
                            === CASE STUDY DOCUMENT ===
                            {case_text[:5000]}
                            
                            === TEACHING NOTE DOCUMENT ===
                            {tn_text[:5000]}
                            
                            Provide your analysis in the following JSON structure:
                            {{
                                "narrative": [a detailed narrative analysis addressing the criterion, as a single string],
                                "document_reference": [specific sections or content from BOTH documents that supports this analysis as a single string]
                            }}
                            
                            {PROMPT_EXCLUSION_INSTRUCTIONS}
                            
                            Important:
                            - Do NOT include a score — this criterion is informational only
                            - Provide a thorough, well-structured narrative
                            - Ensure narrative and document_reference are STRINGS, not lists
                            """
                        elif requires_tn or area_id == "area3":
                            return f"""
                            You are a case study evaluation expert using the CBC-India AGK Case Study Review Rubric.
                            
                            Evaluate the alignment between the case study and its teaching note against this specific criterion:
                            
                            **Criterion:** {criterion_info['name']}
                            **Description:** {criterion_info['description']}
                            **Evaluation Task:** {agent_prompt}
                            **Scoring Guide:** {scoring_logic}
                            **Maximum Score:** {max_score}
                            
                            === CASE STUDY DOCUMENT ===
                            {case_text[:5000]}
                            
                            === TEACHING NOTE DOCUMENT ===
                            {tn_text[:5000]}
                            
                            Provide an analysis with the following JSON structure:
                            {{
                                "score": [a number between 0 and {max_score}, following the scoring guide above],
                                "reasoning": [detailed explanation of why this score was given, based on the scoring logic, as a single string],
                                "document_reference": [specific sections or content from BOTH documents that supports this assessment as a single string]
                            }}
                            
                            {PROMPT_EXCLUSION_INSTRUCTIONS}
                            
                            Important:
                            - The score MUST be an integer between 0 and {max_score}
                            - Follow the scoring logic exactly: {scoring_logic}
                            - Apply the rubric strictly and consistently — use the scoring guide as the sole basis for the score
                            - Do NOT give benefit of the doubt — if evidence is absent or ambiguous, score lower
                            - Be deterministic: the same document should always receive the same score for this criterion
                            - Provide specific evidence from BOTH the case study AND teaching note
                            - Ensure reasoning and document_reference are STRINGS, not lists
                            """
                        else:
                            return f"""
                            You are a case study evaluation expert using the CBC-India AGK Case Study Review Rubric.
                            
                            Evaluate the following case study document against this specific criterion:
                            
                            **Criterion:** {criterion_info['name']}
                            **Description:** {criterion_info['description']}
                            **Evaluation Task:** {agent_prompt}
                            **Scoring Guide:** {scoring_logic}
                            **Maximum Score:** {max_score}
                            
                            Case Study Document:
                            {case_text[:6000]}
                            
                            Provide an analysis with the following JSON structure:
                            {{
                                "score": [a number between 0 and {max_score}, following the scoring guide above],
                                "reasoning": [detailed explanation of why this score was given, based on the scoring logic, as a single string],
                                "document_reference": [specific sections or content from the document that supports this assessment as a single string]
                            }}
                            
                            {PROMPT_EXCLUSION_INSTRUCTIONS}
                            
                            Important:
                            - The score MUST be an integer between 0 and {max_score}
                            - Follow the scoring logic exactly: {scoring_logic}
                            - Apply the rubric strictly and consistently — use the scoring guide as the sole basis for the score
                            - Do NOT give benefit of the doubt — if evidence is absent or ambiguous, score lower
                            - Be deterministic: the same document should always receive the same score for this criterion
                            - Provide specific evidence from the document
                            - Ensure reasoning and document_reference are STRINGS, not lists
                            """

                    def _evaluate_criterion(area_id, criterion_id, criterion_info):
                        prompt = _build_criterion_prompt(area_id, criterion_info)
                        result = call_openai_api(prompt, response_format="json_object", temperature=0.1, seed=42)
                        return area_id, criterion_id, criterion_info, result

                    def _process_result(area_id, criterion_id, criterion_info, result):
                        is_informational = criterion_info.get('informational', False)
                        max_score = criterion_info.get('max_score', 3) if not is_informational else 0
                        if is_informational:
                            if isinstance(result.get("narrative"), list):
                                result["narrative"] = ". ".join(result["narrative"])
                            if not result.get("narrative"):
                                result["narrative"] = result.get("reasoning", "No analysis available.")
                            if isinstance(result.get("document_reference"), list):
                                result["document_reference"] = ". ".join(result["document_reference"])
                            result["score"] = 0
                            result["informational"] = True
                        else:
                            if isinstance(result.get("reasoning"), list):
                                result["reasoning"] = ". ".join(result["reasoning"])
                            if isinstance(result.get("document_reference"), list):
                                result["document_reference"] = ". ".join(result["document_reference"])
                            score = result.get("score", 0)
                            if isinstance(score, (int, float)):
                                result["score"] = min(max(0, int(score)), max_score)
                            else:
                                result["score"] = 0
                        return result

                    for area_id, area_info in ASSESSMENT_AREAS.items():
                        progress_text.text(f"Analyzing: {area_info['name']} (parallel)...")
                        criteria = ASSESSMENT_CRITERIA[area_id]
                        st.session_state.assessment_results[area_id] = {}

                        with ThreadPoolExecutor(max_workers=len(criteria)) as executor:
                            futures = {
                                executor.submit(_evaluate_criterion, area_id, cid, cinfo): (cid, cinfo)
                                for cid, cinfo in criteria.items()
                            }
                            for future in as_completed(futures):
                                try:
                                    a_id, c_id, c_info, raw_result = future.result()
                                    processed = _process_result(a_id, c_id, c_info, raw_result)
                                    st.session_state.assessment_results[a_id][c_id] = processed
                                except Exception as e:
                                    cid_fallback, cinfo_fallback = futures[future]
                                    st.session_state.assessment_results[area_id][cid_fallback] = {
                                        "score": 0, "reasoning": f"Error: {str(e)}", "document_reference": "Error"
                                    }

                        completed_steps += 1
                        progress_bar.progress(completed_steps / total_steps)

                    progress_text.text("Running Writing Quality Analysis, KCM Mapping & Sector Tags (parallel)...")
                    
                    writing_future = None
                    kcm_future = None
                    sector_future = None

                    kcm_prompt = f"""
                    You are an expert in the Karmayogi Competency Model (KCM) used by the Government of India for civil service capacity building.
                    
                    Based on the case study below, identify the TOP 3-4 SUB-COMPETENCIES (not just the broad competencies) that are most relevant to this case.
                    
                    === KARMAYOGI COMPETENCY MODEL (Competencies & their Sub-competencies) ===
                    
                    BEHAVIORAL:
                    - Integrity & Ethics: Honesty, Fairness, Moral Courage, Consistency, Transparency
                    - Adaptability: Flexibility, Openness to Change, Resilience, Versatility
                    - Compassion: Empathy, Sensitivity, Kindness, Supportive
                    - Perpetual Learning: Self-Development, Inquisitiveness, Knowledge Sharing, Reflective Practice
                    - Commitment & Purpose: Dedication, Goal Orientation, Public Service Value, Mission Focus
                    - Inner Calm & Balance: Emotional Intelligence, Stress Management, Equanimity, Self-Regulation
                    - Attention to Detail: Precision, Thoroughness, Meticulousness, Quality Consciousness
                    
                    FUNCTIONAL:
                    - Citizen Centricity: User-Centric Design, Responsive Service, Public Interest, Service Delivery Focus
                    - Accountability: Responsibility, Transparency, Results Orientation, Ethical Governance
                    - Innovation & Technology: Digital Literacy, Creative Problem Solving, Process Improvement, Tech Adoption
                    - Collaboration & Unity: Teamwork, Stakeholder Engagement, Partnership Building, Conflict Resolution
                    - Strategic Thinking: Visionary Planning, Systems Thinking, Policy Analysis, Risk Assessment
                    - Inclusive Development: Social Equity, Diversity & Inclusion, Sustainable Growth, Poverty Alleviation
                    - Cultural Awareness (Garva): Heritage Appreciation, Local Context Awareness, Indigenous Knowledge Support, National Identity
                    - Service Excellence: Standard Setting, Efficiency, Continuous Improvement, Benchmarking
                    
                    === CASE STUDY ===
                    {case_text[:5000]}
                    
                    {PROMPT_EXCLUSION_INSTRUCTIONS}
                    
                    Provide your analysis in the following JSON structure:
                    {{
                        "behavioral_competencies": [
                            {{
                                "name": "sub-competency name",
                                "parent_competency": "broad competency name",
                                "justification": "brief justification of why this sub-competency is relevant"
                            }}
                        ],
                        "functional_competencies": [
                            {{
                                "name": "sub-competency name", 
                                "parent_competency": "broad competency name",
                                "justification": "brief justification of why this sub-competency is relevant"
                            }}
                        ]
                    }}
                    
                    Important:
                    - Identify specific SUB-COMPETENCIES from the provided list.
                    - Select 2 behavioral and 2 functional sub-competencies (total 4).
                    - Ensure all values are strings.
                    """
                    
                    sector_prompt = f"""
                    You are an expert at classifying case studies by sector, sub-theme, and keywords.
                    
                    Based on the case study below, identify the applicable sectors, sub-themes, and keywords.
                    
                    === AVAILABLE SECTORS AND SUB-THEMES ===
                    {json.dumps(SECTOR_MAPPING, indent=2)}
                    
                    === CASE STUDY ===
                    {case_text[:5000]}
                    
                    {PROMPT_EXCLUSION_INSTRUCTIONS}
                    
                    Provide your analysis in the following JSON structure:
                    {{
                        "sectors": ["sector1", "sector2"],
                        "subthemes": {{"sector1": ["subtheme1", "subtheme2"], "sector2": ["subtheme1"]}},
                        "keywords": ["keyword1", "keyword2", "keyword3"]
                    }}
                    
                    Important:
                    - Select ONLY sectors from the provided SECTOR_MAPPING list
                    - For each selected sector, identify relevant sub-themes from that sector's sub-theme list
                    - Generate 5-10 keywords that a general public audience would use to search for this case study
                    - Keywords should be simple, commonly used terms
                    """

                    with ThreadPoolExecutor(max_workers=3) as executor:
                        writing_future = executor.submit(
                            analyze_writing_quality_chunked,
                            case_text,
                            exclusion_instructions=PROMPT_EXCLUSION_INSTRUCTIONS
                        )
                        kcm_future = executor.submit(
                            call_openai_api, kcm_prompt, "gpt-4o", "json_object", 0.1, 42
                        )
                        sector_future = executor.submit(
                            call_openai_api, sector_prompt, "gpt-4o", "json_object", 0.1, 42
                        )

                        try:
                            st.session_state.writing_findings = writing_future.result()
                        except Exception as e:
                            st.session_state.writing_findings = []
                        completed_steps += 1
                        progress_bar.progress(completed_steps / total_steps)
                        progress_text.text("Finalising competency mapping and sector tags...")

                        try:
                            competency_result = kcm_future.result()
                            st.session_state.competency_mapping = competency_result
                        except Exception:
                            st.session_state.competency_mapping = {}
                        completed_steps += 1
                        progress_bar.progress(completed_steps / total_steps)

                        try:
                            tags_result = sector_future.result()
                            if isinstance(tags_result, dict):
                                st.session_state.sector_tags = tags_result.get("sectors", [])
                                st.session_state.sector_subthemes = tags_result.get("subthemes", {})
                                st.session_state.keywords = tags_result.get("keywords", [])
                        except Exception:
                            st.session_state.sector_tags = []
                            st.session_state.sector_subthemes = {}
                            st.session_state.keywords = []
                        completed_steps += 1
                        progress_bar.progress(completed_steps / total_steps)
                    
                    # Calculate weighted scores after assessment
                    st.session_state.weighted_scores = calculate_weighted_score(st.session_state.assessment_results)
                    
                    progress_text.text("Assessment complete!")
                    progress_bar.progress(100)
                    st.success("Case study assessment completed successfully!")
            else:
                # Show which documents are missing
                if not st.session_state.case_study_text and not st.session_state.teaching_note_text:
                    st.warning("Please upload both the Case Study and Teaching Note documents to begin assessment.")
                elif not st.session_state.case_study_text:
                    st.warning("Please upload the Case Study document.")
                else:
                    st.warning("Please upload the Teaching Note document.")
                
        elif st.session_state.sidebar_tab == "History" and st.session_state.get("active_tool") == "analyser":
            st.header("Assessment History")
            
            if st.session_state.user_id:
                try:
                    # Get user's assessment history with error handling
                    with st.spinner("Loading assessment history..."):
                        assessments = get_user_assessments(st.session_state.user_id)
                    
                    # Refresh button for history
                    if st.button("🔄 Refresh History"):
                        st.rerun()
                    
                    if assessments and len(assessments) > 0:
                        st.success(f"Found {len(assessments)} saved assessments.")
                        
                        for assessment in assessments:
                            with st.container():
                                # Simple list of assessments without custom CSS
                                st.subheader(assessment.document_name)
                                st.caption(f"Date: {assessment.created_at.strftime('%Y-%m-%d %H:%M')}")
                                
                                # Simple load button
                                if st.button("Load Assessment", key=f"load_{assessment.id}"):
                                    try:
                                        # Load assessment data into session state with history flag
                                        st.session_state.document_name = assessment.document_name
                                        st.session_state.case_study_summary = assessment.policy_summary
                                        st.session_state.assessment_results = assessment.get_results_dict()
                                        st.session_state.recommendations = assessment.recommendations
                                        st.session_state.case_study_text = "Loaded from history"  # Placeholder for text
                                        
                                        results_dict = st.session_state.assessment_results
                                        st.session_state.sector_tags = results_dict.get('_sector_tags', [])
                                        st.session_state.sector_subthemes = results_dict.get('_sector_subthemes', {})
                                        st.session_state.keywords = results_dict.get('_keywords', [])
                                        
                                        # Recalculate weighted scores
                                        st.session_state.weighted_scores = calculate_weighted_score(st.session_state.assessment_results)
                                        
                                        # Set a special flag to indicate this is from history
                                        st.session_state.loaded_from_history = True
                                        
                                        # Switch to New Assessment tab
                                        st.session_state.sidebar_tab = "New Assessment"
                                        
                                        # Show a success message
                                        st.success(f"Successfully loaded assessment: {assessment.document_name}")
                                        st.rerun()
                                    except Exception as e:
                                        st.error(f"Error loading assessment: {str(e)}")
                                
                                st.markdown("---")
                    else:
                        st.info("No assessment history found. Complete an assessment and save it to see it here.")
                        st.write("After saving an assessment, click the 'Refresh History' button to see it listed here.")
                except Exception as e:
                    st.error(f"Error retrieving assessment history: {str(e)}")
                    st.info("Please try refreshing or check database connection.")
        if st.session_state.get("active_tool") == "caseconnect":
            st.header("Course Outline")
            st.caption("Upload your syllabus, session plan, or curriculum outline (optional but encouraged)")
            curriculum_file = st.file_uploader("Choose file", type=["pdf", "docx"], key="curriculum_uploader")
            
            if curriculum_file is not None:
                try:
                    if curriculum_file.type == "application/pdf":
                        st.session_state.caseconnect_curriculum_text = extract_text_from_pdf(curriculum_file)
                    elif curriculum_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                        st.session_state.caseconnect_curriculum_text = extract_text_from_docx(curriculum_file)
                    st.success(f"Uploaded: {curriculum_file.name}")
                    st.info(f"Extracted {len(st.session_state.caseconnect_curriculum_text)} characters")
                except Exception as e:
                    st.error(f"Error processing file: {str(e)}")
            
            st.markdown("---")
            st.info(f"Referencing {CASE_COUNT} case studies from the AGK repository")

    elif st.session_state.logged_in:
        st.info("Choose a tool from the home page to get started.")
    else:
        st.info("Please login to access the tools.")

# Main content area
if st.session_state.get("active_tool") is None and st.session_state.logged_in:
    st.header("Assessment Framework")
    st.write("The Case Study Analyser uses the CBC-India AGK Case Study Review Rubric with four weighted assessment areas:")
    
    for area_id, area_info in ASSESSMENT_AREAS.items():
        with st.expander(f"{area_info['name']} (Weight: {area_info['weight']*100:.0f}%, Total Points: {area_info['total_points']})"):
            st.write(area_info["description"])
            st.subheader("Criteria:")
            criteria = ASSESSMENT_CRITERIA[area_id]
            for criterion_id, criterion_info in criteria.items():
                if criterion_info.get("informational", False):
                    st.write(f"**{criterion_info['name']}** (Informational — no score)")
                    st.write(f"  - {criterion_info['description']}")
                else:
                    st.write(f"**{criterion_info['name']}** (Max: {criterion_info.get('max_score', 3)} points)")
                    st.write(f"  - {criterion_info['description']}")
                    st.write(f"  - *Scoring:* {criterion_info.get('scoring_logic', '')}")

    st.markdown("---")

    tool_col1, tool_col2, tool_col3 = st.columns(3)

    with tool_col1:
        st.markdown("""
        <div class="card" style="min-height: 320px;">
            <h3 style="color: #1E3A8A; margin-top: 0;">Case Study Analyser</h3>
            <p>Evaluate case studies against the CBC-India AGK Case Study Review Rubric. Upload your case study and teaching note documents, and the AI will assess them across four weighted areas — Structure, Language, Alignment, and Effectiveness — generating a detailed scoring report with recommendations.</p>
        </div>
        """, unsafe_allow_html=True)
        if st.button("Get Started — Case Study Analyser", key="start_analyser", use_container_width=True):
            st.session_state.active_tool = "analyser"
            st.session_state.sidebar_tab = "New Assessment"
            st.rerun()

    with tool_col2:
        st.markdown(f"""
        <div class="card" style="min-height: 320px;">
            <h3 style="color: #1E3A8A; margin-top: 0;">CaseConnect</h3>
            <p>Discover relevant governance case studies from the AGK repository for your teaching programme. Upload your course outline and answer a short questionnaire — the tool will recommend cases aligned with your course design, competencies, and sector. Currently referencing {CASE_COUNT} case studies.</p>
        </div>
        """, unsafe_allow_html=True)
        if st.button("Get Started — CaseConnect", key="start_caseconnect", use_container_width=True):
            st.session_state.active_tool = "caseconnect"
            st.rerun()

    with tool_col3:
        st.markdown("""
        <div class="card" style="min-height: 320px;">
            <h3 style="color: #1E3A8A; margin-top: 0;">Case Study Generator</h3>
            <p>Draft a case study from raw source material — interviews, reports, websites and notes. Choose between Lesson-Drawing, Decision-Forcing or Caselet formats; the tool builds the chronology, stakeholder map and dilemma, drafts each section against the AGK template, runs language and structure checks, and offers an optional teaching note.</p>
        </div>
        """, unsafe_allow_html=True)
        if st.button("Get Started — Case Study Generator", key="start_generator", use_container_width=True):
            st.session_state.active_tool = "generator"
            st.session_state.gen_step = 1
            st.rerun()

elif st.session_state.get("active_tool") == "caseconnect" and st.session_state.logged_in:
    st.header("Course Discovery Questionnaire")
    st.write("Answer the following questions to help us recommend the most relevant case studies for your teaching programme.")

    all_competency_names = []
    for category in ["behavioral", "functional"]:
        for comp_key, comp_data in KCM_COMPETENCIES.get(category, {}).items():
            all_competency_names.append(comp_data["name"])

    all_sectors = list(SECTOR_MAPPING.keys())

    with st.form("caseconnect_form"):
        q_learners = st.text_area(
            "1. Who are the learners?",
            placeholder="e.g., Mid-career IAS officers, State government training faculty, Fresh recruits to civil services...",
            height=80
        )
        
        q_objective = st.text_area(
            "2. What is the primary learning objective?",
            placeholder="e.g., Understanding stakeholder coordination in large infrastructure projects, Developing leadership in crisis situations...",
            height=80
        )
        
        q_competencies = st.multiselect(
            "3. Which competencies are you focusing on? (Karmayogi Competency Model)",
            options=all_competency_names,
            default=[]
        )
        
        q_duration = st.selectbox(
            "4. Duration of the session?",
            options=["30 minutes", "60 minutes", "90 minutes", "2 hours", "Half day (3-4 hours)", "Full day", "Multi-day programme"]
        )
        
        q_sector = st.multiselect(
            "5. What is the Sector/Theme?",
            options=all_sectors,
            default=[]
        )

        submitted = st.form_submit_button("Find Matching Cases", use_container_width=True)

    if submitted:
        if not q_learners and not q_objective and not q_competencies and not q_sector:
            st.warning("Please fill in at least one question to get recommendations.")
        else:
            questionnaire = {
                "learners": q_learners,
                "objective": q_objective,
                "competencies": q_competencies,
                "duration": q_duration,
                "sector": ", ".join(q_sector) if q_sector else ""
            }
            
            with st.spinner("Analysing your inputs and finding matching cases..."):
                st.session_state.caseconnect_results = run_caseconnect_analysis(
                    questionnaire,
                    st.session_state.caseconnect_curriculum_text,
                    CASE_DATABASE_TEXT,
                    CASE_COUNT
                )

    if st.session_state.caseconnect_results:
        results = st.session_state.caseconnect_results
        
        st.markdown("---")
        st.header("Recommended Cases for Your Course")
        
        case_recs = results.get("case_recommendations", [])
        if case_recs:
            import html as html_mod
            for i, rec in enumerate(case_recs, 1):
                with st.expander(f"**{i}. {rec.get('case_title', 'Untitled')}**", expanded=(i <= 3)):
                    st.write(f"**Why it fits:** {rec.get('why_it_fits', '')}")

                    themes = rec.get("key_themes", [])
                    if isinstance(themes, str):
                        themes = [themes]
                    if themes:
                        theme_badges = " ".join([f'<span class="score-badge" style="background-color:#EFF6FF;color:#1E3A8A;border:1px solid #BFDBFE;">{html_mod.escape(str(t))}</span>' for t in themes])
                        st.markdown(f"**Key Themes:** {theme_badges}", unsafe_allow_html=True)

                    comps = rec.get("relevant_competencies", [])
                    if isinstance(comps, str):
                        comps = [comps]
                    if comps:
                        comp_badges = " ".join([f'<span class="score-badge score-good">{html_mod.escape(str(c))}</span>' for c in comps])
                        st.markdown(f"**Relevant Competencies:** {comp_badges}", unsafe_allow_html=True)

                    discussion = rec.get("discussion_points", [])
                    if isinstance(discussion, str):
                        discussion = [discussion]
                    if discussion:
                        st.write("**Discussion Points for Classroom:**")
                        for dp_idx, dp in enumerate(discussion, 1):
                            st.write(f"{dp_idx}. {dp}")

                    dur = rec.get("suggested_duration", "")
                    if dur:
                        st.write(f"**Suggested Duration:** {dur}")

                    igot_link = _resolve_igot_link(rec.get("case_title", ""))
                    if igot_link:
                        safe_link = html_mod.escape(igot_link)
                        st.markdown(
                            f'<a href="{safe_link}" target="_blank" style="display:inline-block;background-color:#074fa5;color:white;padding:8px 20px;border-radius:6px;text-decoration:none;font-weight:600;font-size:14px;margin-top:8px;">Go to Case Study on iGOT ↗</a>',
                            unsafe_allow_html=True
                        )
        else:
            st.info("No specific case recommendations were generated. Try providing more details in the questionnaire.")

        module_suggestions = results.get("module_suggestions", [])
        if module_suggestions:
            st.markdown("---")
            st.header("Module-Level Suggestions")
            for mod in module_suggestions:
                st.subheader(mod.get("module_name", "Module"))
                mod_cases = mod.get("recommended_cases", [])
                if mod_cases:
                    for mc in mod_cases:
                        st.write(f"- **{mc.get('case_title', '')}** — {mc.get('relevance', '')}")

        teaching_strategy = results.get("teaching_strategy", "")
        if teaching_strategy:
            st.markdown("---")
            st.header("Teaching Strategy")
            st.write(teaching_strategy)

        additional = results.get("additional_notes", "")
        if additional:
            st.markdown("---")
            st.header("Additional Notes")
            st.write(additional)

elif st.session_state.get("active_tool") == "generator" and st.session_state.logged_in:
    # ------------------------------------------------------------------
    # Case Study Generator — multi-step wizard
    # ------------------------------------------------------------------
    GEN_STEP_LABELS = [
        "1. Choose case type",
        "2. Case metadata",
        "3. Upload sources",
        "4. Narrative intent",
        "5. Source processing",
        "6. Draft sections",
        "7. Compliance review",
        "8. Teaching note & export",
    ]
    st.markdown(
        " · ".join(
            f"**{lbl}**" if (i + 1) == st.session_state.gen_step else lbl
            for i, lbl in enumerate(GEN_STEP_LABELS)
        )
    )
    st.markdown("---")

    def _gen_persist():
        """Persist the in-progress draft to the database (best-effort)."""
        if not st.session_state.user_id:
            return
        try:
            cid = save_generated_case(
                user_id=st.session_state.user_id,
                case_id=st.session_state.gen_case_id,
                title=(st.session_state.gen_metadata or {}).get("title") or "Untitled",
                case_type=st.session_state.gen_case_type,
                metadata=st.session_state.gen_metadata,
                intent=st.session_state.gen_intent,
                sources=st.session_state.gen_sources,
                processed=st.session_state.gen_processed,
                sections=st.session_state.gen_sections,
                compliance=st.session_state.gen_compliance,
                teaching_note=st.session_state.gen_teaching_note,
                kcm_mapping=st.session_state.gen_kcm_mapping,
                final_documents=st.session_state.gen_final_documents,
            )
            if cid:
                st.session_state.gen_case_id = cid
        except Exception as e:
            st.warning(f"Could not auto-save draft: {e}")

    def _gen_nav(prev_to=None, next_to=None, next_label="Next →"):
        cols = st.columns([1, 1, 4])
        with cols[0]:
            if prev_to is not None and st.button("← Back", key=f"gen_back_{st.session_state.gen_step}"):
                st.session_state.gen_step = prev_to
                st.rerun()
        with cols[1]:
            if next_to is not None and st.button(next_label, key=f"gen_next_{st.session_state.gen_step}"):
                _gen_persist()
                st.session_state.gen_step = next_to
                st.rerun()

    # ----- Step 1: case type ---------------------------------------------
    if st.session_state.gen_step == 1:
        st.header("Step 1 — Choose the type of case study")
        st.write("Select the format that best matches what you want to write. You can change later, but the structure and word-count rules differ.")

        for ctype, meta in CASE_TYPES.items():
            with st.container():
                st.subheader(meta["label"])
                st.write(meta["description"])
                st.caption(f"Word count: {meta['word_min']:,} – {meta['word_max']:,}")
                if st.button(f"Choose {meta['label']}", key=f"choose_{ctype}"):
                    st.session_state.gen_case_type = ctype
                    if not st.session_state.gen_metadata.get("word_target"):
                        st.session_state.gen_metadata["word_target"] = meta["default_target"]
                    st.session_state.gen_step = 2
                    st.rerun()
                st.markdown("---")

    # ----- Step 2: metadata ----------------------------------------------
    elif st.session_state.gen_step == 2:
        st.header("Step 2 — Case metadata")
        st.write("Provide the basic facts that anchor the case.")
        ctype = st.session_state.gen_case_type
        meta = st.session_state.gen_metadata or {}
        ctype_meta = CASE_TYPES[ctype]

        with st.form("gen_metadata_form"):
            title = st.text_input("Case title", value=meta.get("title", ""))
            authors = st.text_input("Author(s)", value=meta.get("authors", ""))
            initiative = st.text_input("Initiative or programme name", value=meta.get("initiative", ""))
            region = st.text_input("Country / state / region", value=meta.get("region", ""))
            sector = st.text_input("Sector or theme", value=meta.get("sector", ""))
            protagonist = st.text_input("Primary protagonist (name and role)", value=meta.get("protagonist", ""))
            stakeholders = st.text_area(
                "Other key stakeholders (one per line or comma-separated)",
                value=meta.get("stakeholders", ""),
                height=80,
            )
            time_period = st.text_input("Time period covered (e.g., 2018-2024)", value=meta.get("time_period", ""))
            audience = st.text_input("Intended teaching audience", value=meta.get("audience", ""))
            word_target = st.number_input(
                f"Target word count (recommended {ctype_meta['word_min']}–{ctype_meta['word_max']})",
                min_value=int(ctype_meta["word_min"] * 0.5),
                max_value=int(ctype_meta["word_max"] * 1.5),
                value=int(meta.get("word_target") or ctype_meta["default_target"]),
                step=100,
            )
            competencies = st.text_input(
                "Competency to align with (optional)",
                value=meta.get("competencies", ""),
            )
            note_from_author = st.text_area(
                "Note from the Author (appears as front matter)",
                value=meta.get("note_from_author", ""),
                height=80,
            )
            acknowledgements = st.text_area(
                "Acknowledgements (optional — appears as front matter)",
                value=meta.get("acknowledgements", ""),
                height=80,
            )
            teaching_note = st.checkbox(
                "Generate a teaching note alongside the case",
                value=bool(meta.get("teaching_note", True)),
            )
            saved = st.form_submit_button("Save and continue →")

        if saved:
            mandatory = [
                ("Case title", title),
                ("Author(s)", authors),
                ("Initiative or programme name", initiative),
                ("Country / state / region", region),
                ("Sector or theme", sector),
                ("Primary protagonist (name and role)", protagonist),
                ("Time period covered (e.g., 2018-2024)", time_period),
                ("Note from the Author (appears as front matter)", note_from_author),
            ]
            missing = [label for label, val in mandatory if not (val or "").strip()]
            target_word_label = (
                f"Target word count (recommended {ctype_meta['word_min']}–{ctype_meta['word_max']})"
            )
            try:
                if int(word_target) <= 0:
                    missing.append(target_word_label)
            except (TypeError, ValueError):
                missing.append(target_word_label)
            if missing:
                st.error(
                    "Please fill the following required fields before continuing: "
                    + ", ".join(missing)
                )
                st.stop()

            st.session_state.gen_metadata = {
                "title": title.strip(),
                "authors": authors.strip(),
                "initiative": initiative.strip(),
                "region": region.strip(),
                "sector": sector.strip(),
                "protagonist": protagonist.strip(),
                "stakeholders": stakeholders.strip(),
                "time_period": time_period.strip(),
                "audience": audience.strip(),
                "word_target": int(word_target),
                "competencies": competencies.strip(),
                "note_from_author": note_from_author.strip(),
                "acknowledgements": acknowledgements.strip(),
                "teaching_note": bool(teaching_note),
            }
            _gen_persist()
            st.session_state.gen_step = 3
            st.rerun()

        _gen_nav(prev_to=1, next_to=None)

    # ----- Step 3: source upload -----------------------------------------
    elif st.session_state.gen_step == 3:
        st.header("Step 3 — Upload source material")
        st.write(
            "Upload every source you want the AI to draw on: interview transcripts, reports, "
            "news articles, briefs and your own author notes. The tool only writes from these sources."
        )

        uploaded = st.file_uploader(
            "Upload files (PDF, DOCX or TXT). You may upload several.",
            type=["pdf", "docx", "txt"],
            accept_multiple_files=True,
            key="gen_source_uploader",
        )

        if uploaded:
            existing_names = {s["name"] for s in st.session_state.gen_sources}
            for f in uploaded:
                if f.name in existing_names:
                    continue
                try:
                    fname_lower = f.name.lower()
                    if f.type == "application/pdf" or fname_lower.endswith(".pdf"):
                        text = extract_text_from_pdf(f)
                    elif fname_lower.endswith(".docx"):
                        text = extract_text_from_docx(f)
                    else:
                        raw = f.read()
                        if isinstance(raw, bytes):
                            text = raw.decode("utf-8", errors="replace")
                        else:
                            text = str(raw)
                    # Apply the repair pass used elsewhere in the suite to fix
                    # PDF/DOCX extraction artefacts (broken words, ligature
                    # splits, etc.) before sources reach the AI prompts.
                    try:
                        text = repair_broken_words(text)
                    except Exception:
                        pass
                except Exception as e:
                    st.error(f"Could not read {f.name}: {e}")
                    continue
                st.session_state.gen_sources.append({
                    "name": f.name,
                    "type": "Uploaded document",
                    "speaker": "",
                    "date": "",
                    "text": text,
                })
            _gen_persist()

        with st.expander("Add sources by URL (one per line)"):
            with st.form("gen_url_sources", clear_on_submit=True):
                url_block = st.text_area(
                    "Paste one URL per line. The page text will be fetched and added as a source.",
                    placeholder="https://example.gov.in/report\nhttps://news.example.com/article",
                    height=120,
                )
                fetch_clicked = st.form_submit_button("Fetch URLs and add as sources")
            if fetch_clicked and url_block.strip():
                import socket
                import ipaddress

                def _is_safe_public_url(u):
                    """Block SSRF: reject non-public IP targets, metadata endpoints,
                    non-default ports and non-http(s) schemes."""
                    try:
                        p = urlparse(u)
                    except Exception:
                        return False, "Could not parse URL"
                    if p.scheme not in ("http", "https"):
                        return False, "Only http(s) URLs are accepted"
                    if not p.hostname:
                        return False, "Missing hostname"
                    # Disallow non-standard ports to reduce internal-service probing surface.
                    if p.port not in (None, 80, 443):
                        return False, f"Port {p.port} is not allowed"
                    host = p.hostname.lower()
                    if host in ("localhost", "metadata.google.internal", "metadata"):
                        return False, "Host is blocked"
                    try:
                        infos = socket.getaddrinfo(host, None)
                    except socket.gaierror:
                        return False, "DNS lookup failed"
                    for info in infos:
                        ip = info[4][0]
                        try:
                            ip_obj = ipaddress.ip_address(ip)
                        except ValueError:
                            return False, f"Resolved to non-IP {ip}"
                        if (
                            ip_obj.is_private
                            or ip_obj.is_loopback
                            or ip_obj.is_link_local
                            or ip_obj.is_reserved
                            or ip_obj.is_multicast
                            or ip_obj.is_unspecified
                        ):
                            return False, f"Resolved to non-public address {ip}"
                    return True, None

                # Custom handler that REFUSES to follow redirects, so a public
                # initial host cannot be used to bounce into a private network.
                class _NoRedirect(urllib.request.HTTPRedirectHandler):
                    def http_error_301(self, req, fp, code, msg, headers):
                        raise urllib.error.HTTPError(req.full_url, code, "Redirects are blocked", headers, fp)
                    http_error_302 = http_error_301
                    http_error_303 = http_error_301
                    http_error_307 = http_error_301
                    http_error_308 = http_error_301

                _opener = urllib.request.build_opener(_NoRedirect())

                MAX_BYTES = 2_000_000  # 2 MB cap per page
                urls = [u.strip() for u in url_block.splitlines() if u.strip()]
                fetched = 0
                for url in urls:
                    safe, reason = _is_safe_public_url(url)
                    if not safe:
                        st.error(f"Skipped {url}: {reason}")
                        continue
                    parsed = urlparse(url)
                    try:
                        req = urllib.request.Request(
                            url,
                            headers={"User-Agent": "Mozilla/5.0 (compatible; CBC-AGK-Generator/1.0)"},
                        )
                        with _opener.open(req, timeout=15) as resp:
                            charset = resp.headers.get_content_charset() or "utf-8"
                            raw = resp.read(MAX_BYTES + 1).decode(charset, errors="replace")
                            if len(raw.encode(charset, errors="replace")) > MAX_BYTES:
                                st.warning(f"{url} was larger than 2 MB — only the first 2 MB was kept.")
                        # Cheap HTML-to-text: drop scripts/styles, strip tags.
                        cleaned = re.sub(r"<script[^>]*>.*?</script>", " ", raw, flags=re.S | re.I)
                        cleaned = re.sub(r"<style[^>]*>.*?</style>", " ", cleaned, flags=re.S | re.I)
                        cleaned = re.sub(r"<[^>]+>", " ", cleaned)
                        cleaned = re.sub(r"\s+", " ", cleaned).strip()
                        try:
                            cleaned = repair_broken_words(cleaned)
                        except Exception:
                            pass
                        if not cleaned:
                            st.error(f"Fetched {url} but the page contained no extractable text.")
                            continue
                        st.session_state.gen_sources.append({
                            "name": parsed.netloc + parsed.path[:60],
                            "type": "Website / URL",
                            "speaker": parsed.netloc,
                            "date": "",
                            "text": f"[Source URL: {url}]\n\n{cleaned}",
                        })
                        fetched += 1
                    except urllib.error.HTTPError as e:
                        if "Redirects are blocked" in str(e):
                            st.error(
                                f"Skipped {url}: server tried to redirect (status {e.code}); "
                                "redirects are disabled to prevent SSRF. Use the final URL directly."
                            )
                        else:
                            st.error(f"Could not fetch {url}: HTTP {e.code} — {e.reason}")
                    except Exception as e:
                        st.error(f"Could not fetch {url}: {e}")
                if fetched:
                    _gen_persist()
                    st.success(f"Added {fetched} URL source(s).")
                    st.rerun()

        with st.expander("Add a source by pasting text (URL content, raw notes, transcripts)"):
            with st.form("gen_paste_source", clear_on_submit=True):
                ps_name = st.text_input("Source label", placeholder="e.g., Interview with VPA Chairman, March 2024")
                ps_type = st.selectbox(
                    "Source type",
                    ["Interview transcript", "Report", "News article", "Website / URL", "Author notes", "Other"],
                )
                ps_speaker = st.text_input("Speaker / author / publication")
                ps_date = st.text_input("Date (as known)")
                ps_text = st.text_area("Paste source text", height=200)
                added = st.form_submit_button("Add source")
            if added and ps_text.strip():
                st.session_state.gen_sources.append({
                    "name": ps_name or f"Source {len(st.session_state.gen_sources)+1}",
                    "type": ps_type,
                    "speaker": ps_speaker,
                    "date": ps_date,
                    "text": ps_text.strip(),
                })
                _gen_persist()
                st.success("Source added.")
                st.rerun()

        if st.session_state.gen_sources:
            st.subheader("Source bundle")
            for i, s in enumerate(st.session_state.gen_sources):
                with st.expander(f"Source {i+1}: {s['name']}  ·  {len(s.get('text','')):,} chars"):
                    cols = st.columns(4)
                    with cols[0]:
                        s["type"] = st.text_input("Type", value=s.get("type", ""), key=f"st_type_{i}")
                    with cols[1]:
                        s["speaker"] = st.text_input("Speaker / author", value=s.get("speaker", ""), key=f"st_sp_{i}")
                    with cols[2]:
                        s["date"] = st.text_input("Date", value=s.get("date", ""), key=f"st_dt_{i}")
                    with cols[3]:
                        if st.button("Remove", key=f"st_rm_{i}"):
                            st.session_state.gen_sources.pop(i)
                            _gen_persist()
                            st.rerun()
                    st.text_area("Preview", value=s.get("text", "")[:3000], height=180, key=f"st_prev_{i}", disabled=True)

        else:
            st.info("Add at least one source before continuing.")

        _gen_nav(prev_to=2, next_to=4 if st.session_state.gen_sources else None)

    # ----- Step 4: narrative intent --------------------------------------
    elif st.session_state.gen_step == 4:
        st.header("Step 4 — Tell us what story you want to tell")
        st.write(
            "These answers shape the dilemma statement and the narrative arc. Be specific. "
            "Anything you write here is treated as author guidance, not as factual claims."
        )
        intent = st.session_state.gen_intent or {}

        with st.form("gen_intent_form"):
            central = st.text_area("What is the central problem the case explores?", value=intent.get("central_problem", ""), height=80)
            dilemma = st.text_area("What is the dilemma or tension at the heart of the case?", value=intent.get("dilemma", ""), height=80)
            went_well = st.text_area("What went well?", value=intent.get("went_well", ""), height=80)
            did_not = st.text_area("What did not go as expected?", value=intent.get("did_not_go", ""), height=80)
            key_lesson = st.text_area("If the reader takes away one lesson, what should it be?", value=intent.get("key_lesson", ""), height=80)
            opening = st.text_area("Preferred opening scene or hook", value=intent.get("opening_scene", ""), height=80)
            saved = st.form_submit_button("Save and continue →")

        if saved:
            intent_required = [
                ("Central problem", central),
                ("Dilemma or tension", dilemma),
                ("What went well", went_well),
                ("What did not go as expected", did_not),
                ("Key lesson", key_lesson),
                ("Preferred opening scene or hook", opening),
            ]
            missing = [label for label, val in intent_required if not (val or "").strip()]
            if missing:
                st.error(
                    "Please fill the following required fields before continuing: "
                    + ", ".join(missing)
                )
                st.stop()

            st.session_state.gen_intent = {
                "central_problem": central.strip(),
                "dilemma": dilemma.strip(),
                "went_well": went_well.strip(),
                "did_not_go": did_not.strip(),
                "key_lesson": key_lesson.strip(),
                "opening_scene": opening.strip(),
            }
            _gen_persist()
            st.session_state.gen_step = 5
            st.rerun()

        _gen_nav(prev_to=3, next_to=None)

    # ----- Step 5: source processing -------------------------------------
    elif st.session_state.gen_step == 5:
        st.header("Step 5 — Source processing")
        st.write(
            "The tool reads your sources and builds an inventory of what each one contains, a "
            "chronological timeline, a stakeholder map and a draft of the central dilemma. "
            "Review and edit before drafting begins."
        )

        if not st.session_state.gen_processed or st.button("Re-run source processing", key="gen_reprocess"):
            with st.spinner("Reading sources, building chronology, stakeholder map and dilemma..."):
                try:
                    st.session_state.gen_processed = run_source_processing(
                        st.session_state.gen_sources,
                        st.session_state.gen_metadata,
                        st.session_state.gen_intent,
                        st.session_state.gen_case_type,
                    )
                    _gen_persist()
                except GeneratorAPIError as e:
                    st.error(f"Source processing failed: {e}")
                    st.stop()
                except Exception as e:
                    st.error(f"Source processing failed unexpectedly: {e}")
                    st.stop()

        proc = st.session_state.gen_processed or {}

        with st.expander("Source inventory — review and edit before drafting", expanded=False):
            inv_items = (proc.get("inventory") or {}).get("inventory", [])
            if inv_items:
                inv_df_rows = [
                    {
                        "Source": item.get("source_index", idx + 1),
                        "Type": item.get("source_type", ""),
                        "Author / speaker": item.get("speaker_or_author", ""),
                        "Key facts (one per line)": "\n".join(item.get("key_facts", []) or []),
                        "Direct quotes": len(item.get("direct_quotes", []) or []),
                        "Events extracted": len(item.get("events", []) or []),
                    }
                    for idx, item in enumerate(inv_items)
                ]
                inv_edit = st.data_editor(
                    pd.DataFrame(inv_df_rows),
                    use_container_width=True,
                    hide_index=True,
                    num_rows="fixed",
                    disabled=["Source", "Direct quotes", "Events extracted"],
                    key="gen_inventory_editor",
                )
                inv_cols = st.columns([1, 1, 4])
                with inv_cols[0]:
                    if st.button("Save inventory edits", key="gen_save_inventory"):
                        rows = inv_edit.to_dict(orient="records")
                        for idx, row in enumerate(rows):
                            if idx >= len(inv_items):
                                break
                            inv_items[idx]["source_type"] = (row.get("Type") or "").strip()
                            inv_items[idx]["speaker_or_author"] = (row.get("Author / speaker") or "").strip()
                            facts_raw = (row.get("Key facts (one per line)") or "")
                            inv_items[idx]["key_facts"] = [
                                line.strip() for line in facts_raw.splitlines() if line.strip()
                            ]
                        proc.setdefault("inventory", {})["inventory"] = inv_items
                        st.session_state.gen_processed = proc
                        _gen_persist()
                        st.success("Inventory updated.")
                with inv_cols[1]:
                    st.checkbox(
                        "I have reviewed the inventory",
                        key="gen_inventory_confirmed",
                        value=st.session_state.get("gen_inventory_confirmed", False),
                    )
            else:
                st.info("No inventory generated.")

        with st.expander("Chronology — edit before drafting", expanded=True):
            chrono = (proc.get("chronology") or {}).get("chronology", [])
            df_rows = [
                {
                    "Date": c.get("approx_date", ""),
                    "Event": c.get("event", ""),
                    "Actor": c.get("actor", ""),
                    "Outcome": c.get("outcome", ""),
                    "Sources": ", ".join(str(s) for s in (c.get("source_index") or [])),
                }
                for c in chrono
            ] or [{"Date": "", "Event": "", "Actor": "", "Outcome": "", "Sources": ""}]
            chrono_df = st.data_editor(
                pd.DataFrame(df_rows),
                use_container_width=True,
                hide_index=True,
                num_rows="dynamic",
                key="gen_chrono_editor",
            )
            if st.button("Save chronology edits", key="save_chrono_edits"):
                rows = []
                for _, r in chrono_df.iterrows():
                    if not (str(r.get("Event", "")).strip() or str(r.get("Date", "")).strip()):
                        continue
                    src_str = str(r.get("Sources", "") or "")
                    src_indexes = [s.strip() for s in src_str.split(",") if s.strip()]
                    rows.append({
                        "approx_date": str(r.get("Date", "") or "").strip(),
                        "event": str(r.get("Event", "") or "").strip(),
                        "actor": str(r.get("Actor", "") or "").strip(),
                        "outcome": str(r.get("Outcome", "") or "").strip(),
                        "source_index": src_indexes,
                    })
                proc.setdefault("chronology", {})["chronology"] = rows
                st.session_state.gen_processed = proc
                _gen_persist()
                st.success("Chronology saved.")

            gaps = (proc.get("chronology") or {}).get("gaps", [])
            if gaps:
                st.warning("Gaps in the source material:")
                for g in gaps:
                    st.write(f"- {g}")

        with st.expander("Stakeholder map — edit before drafting", expanded=False):
            sks = (proc.get("stakeholders") or {}).get("stakeholders", [])
            df_rows = [
                {
                    "Name": s.get("name", ""),
                    "Designation": s.get("designation", ""),
                    "Side": s.get("side", ""),
                    "Perspective": s.get("perspective", ""),
                    "Tensions": s.get("tensions", ""),
                }
                for s in sks
            ] or [{"Name": "", "Designation": "", "Side": "", "Perspective": "", "Tensions": ""}]
            sks_df = st.data_editor(
                pd.DataFrame(df_rows),
                use_container_width=True,
                hide_index=True,
                num_rows="dynamic",
                key="gen_stake_editor",
            )
            if st.button("Save stakeholder edits", key="save_stake_edits"):
                rows = []
                for _, r in sks_df.iterrows():
                    if not str(r.get("Name", "")).strip():
                        continue
                    rows.append({
                        "name": str(r.get("Name", "") or "").strip(),
                        "designation": str(r.get("Designation", "") or "").strip(),
                        "side": str(r.get("Side", "") or "").strip(),
                        "perspective": str(r.get("Perspective", "") or "").strip(),
                        "tensions": str(r.get("Tensions", "") or "").strip(),
                    })
                proc.setdefault("stakeholders", {})["stakeholders"] = rows
                st.session_state.gen_processed = proc
                _gen_persist()
                st.success("Stakeholder map saved.")

        st.subheader("Central dilemma — review and edit")
        existing = (proc.get("dilemma") or {}).get("dilemma_statement", "")
        edited_dilemma = st.text_area(
            "Dilemma statement",
            value=existing,
            height=160,
            help="This statement will anchor the entire case. Edit freely.",
        )
        if edited_dilemma != existing:
            proc.setdefault("dilemma", {})["dilemma_statement"] = edited_dilemma
            st.session_state.gen_processed = proc
            _gen_persist()

        _gen_nav(prev_to=4, next_to=6, next_label="Continue to drafting →")

    # ----- Step 6: draft sections ----------------------------------------
    elif st.session_state.gen_step == 6:
        ctype = st.session_state.gen_case_type
        st.header("Step 6 — Draft the case")
        st.write(
            "Generate each section in order. After every section is generated, scroll down and "
            "review it. Use 'Regenerate' with feedback to refine."
        )

        sections_meta = CASE_TYPES[ctype]["sections"]
        if ctype == "caselet":
            st.subheader("Caselet draft")
            existing = (st.session_state.gen_sections.get("caselet") or {}).get("text", "")
            feedback = st.text_input("Optional feedback for regeneration", key="caselet_fb")
            if st.button("Generate / regenerate caselet", key="gen_caselet_btn"):
                with st.spinner("Drafting caselet..."):
                    try:
                        text = draft_caselet(
                            st.session_state.gen_metadata,
                            st.session_state.gen_intent,
                            st.session_state.gen_sources,
                            st.session_state.gen_processed,
                            author_feedback=feedback,
                        )
                        st.session_state.gen_sections["caselet"] = {"text": text}
                        _gen_persist()
                        st.rerun()
                    except GeneratorAPIError as e:
                        st.error(f"Caselet drafting failed: {e}")
                    except Exception as e:
                        st.error(f"Caselet drafting failed unexpectedly: {e}")
            if existing:
                st.text_area("Caselet text (editable)", value=existing, height=500, key="caselet_edit_box")
                if st.button("Save edits", key="caselet_save_edits"):
                    st.session_state.gen_sections["caselet"] = {"text": st.session_state.caselet_edit_box}
                    _gen_persist()
                    st.success("Saved.")
        else:
            for sid, sname in sections_meta:
                with st.expander(f"{sname}", expanded=not st.session_state.gen_sections.get(sid)):
                    sec = st.session_state.gen_sections.get(sid, {})
                    existing = sec.get("text", "")
                    fb_key = f"fb_{sid}"
                    feedback = st.text_input("Optional feedback for regeneration", key=fb_key, value=sec.get("feedback", ""))
                    if st.button(("Regenerate " if existing else "Generate ") + sname, key=f"btn_{sid}"):
                        with st.spinner(f"Drafting {sname}..."):
                            try:
                                text = draft_case_section(
                                    ctype,
                                    sid,
                                    st.session_state.gen_metadata,
                                    st.session_state.gen_intent,
                                    st.session_state.gen_sources,
                                    st.session_state.gen_processed,
                                    previous_sections=st.session_state.gen_sections,
                                    author_feedback=feedback,
                                )
                            except Exception as e:
                                st.error(f"Drafting failed: {e}")
                                text = existing
                        st.session_state.gen_sections[sid] = {"text": text, "feedback": feedback}
                        _gen_persist()
                        st.rerun()
                    if existing:
                        word_count = len(re.findall(r"\b\w+\b", existing))
                        st.caption(f"Word count: {word_count}")
                        edit_key = f"edit_{sid}"
                        st.text_area("Section text (editable)", value=existing, height=320, key=edit_key)
                        if st.button("Save edits", key=f"save_{sid}"):
                            st.session_state.gen_sections[sid] = {
                                "text": st.session_state[edit_key],
                                "feedback": feedback,
                            }
                            _gen_persist()
                            st.success("Saved.")

        # Gate progression: every REQUIRED section must have non-empty text.
        # Sections whose label includes "(Optional)" — currently the
        # Decision-Forcing epilogue — do not block progression.
        def _is_optional(sname):
            return "(optional)" in (sname or "").lower()

        required_sids = [sid for sid, sname in sections_meta if not _is_optional(sname)]
        sections_complete = all(
            (st.session_state.gen_sections.get(sid) or {}).get("text", "").strip()
            for sid in required_sids
        )
        if not sections_complete:
            missing = [
                sname for sid, sname in sections_meta
                if not _is_optional(sname)
                and not (st.session_state.gen_sections.get(sid) or {}).get("text", "").strip()
            ]
            st.warning("Generate every required section before continuing. Outstanding: " + ", ".join(missing))
        _gen_nav(prev_to=5, next_to=7 if sections_complete else None,
                 next_label="Continue to compliance review →")

    # ----- Step 7: compliance review -------------------------------------
    elif st.session_state.gen_step == 7:
        st.header("Step 7 — Compliance review")
        st.write(
            "Run automated checks on the full draft for British English, tense, adjectives, em-dashes, "
            "neutrality, structure, factual traceability, abbreviations and word count."
        )

        if not st.session_state.gen_compliance or st.button("Re-run compliance checks", key="rerun_compliance"):
            with st.spinner("Running compliance passes..."):
                try:
                    st.session_state.gen_compliance = run_compliance_passes(
                        st.session_state.gen_case_type,
                        st.session_state.gen_sections,
                        st.session_state.gen_metadata,
                        st.session_state.gen_sources,
                    )
                    _gen_persist()
                except GeneratorAPIError as e:
                    st.error(f"Compliance review failed: {e}")
                    st.stop()
                except Exception as e:
                    st.error(f"Compliance review failed unexpectedly: {e}")
                    st.stop()

        comp = st.session_state.gen_compliance or {}
        m_cols = st.columns(3)
        m_cols[0].metric("Word count", f"{comp.get('word_count', 0):,}")
        m_cols[1].metric("Target", f"{comp.get('target', 0):,}")
        findings = comp.get("findings", [])
        m_cols[2].metric("Findings", len(findings))

        if findings:
            st.subheader("Findings — tick the rows whose fixes you want to apply")
            df = pd.DataFrame([
                {
                    "Apply": True,
                    "Type": f.get("type", ""),
                    "Section": f.get("section", ""),
                    "Severity": f.get("severity", ""),
                    "Original": f.get("original_text", ""),
                    "Suggestion": f.get("suggestion", ""),
                    "Explanation": f.get("explanation", ""),
                }
                for f in findings
            ])
            edited = st.data_editor(
                df,
                use_container_width=True,
                hide_index=True,
                disabled=["Type", "Section", "Severity", "Original", "Suggestion", "Explanation"],
                key="gen_findings_editor",
                height=min(500, 60 + len(df) * 35),
            )

            apply_cols = st.columns([1, 1, 3])
            with apply_cols[0]:
                if st.button("Apply selected fixes", key="apply_compliance_fixes_btn"):
                    selected_indices = [
                        i for i, row in edited.iterrows() if bool(row.get("Apply"))
                    ]
                    if not selected_indices:
                        st.warning("No fixes were selected.")
                    else:
                        try:
                            with st.spinner(f"Applying {len(selected_indices)} fix(es) to your draft…"):
                                updated, applied_count, skipped = apply_compliance_fixes(
                                    st.session_state.gen_case_type,
                                    st.session_state.gen_sections,
                                    findings,
                                    selected_indices,
                                )
                            st.session_state.gen_sections = updated
                            # Re-run compliance so the badges/word count refresh.
                            try:
                                st.session_state.gen_compliance = run_compliance_passes(
                                    st.session_state.gen_case_type,
                                    st.session_state.gen_sections,
                                    st.session_state.gen_metadata,
                                    st.session_state.gen_sources,
                                )
                            except GeneratorAPIError as e:
                                st.warning(f"Re-running compliance failed: {e}")
                            except Exception as e:
                                st.warning(f"Re-running compliance failed unexpectedly: {e}")
                            _gen_persist()
                            msg = f"Applied {applied_count} fix(es)."
                            if skipped:
                                msg += f" {len(skipped)} could not be applied (original phrase not found verbatim — edit manually)."
                            st.success(msg)
                            st.rerun()
                        except GeneratorAPIError as e:
                            st.error(f"Could not apply fixes: {e}")
            with apply_cols[1]:
                if st.button("Re-run compliance", key="rerun_compliance_btn"):
                    try:
                        with st.spinner("Re-running compliance review…"):
                            st.session_state.gen_compliance = run_compliance_passes(
                                st.session_state.gen_case_type,
                                st.session_state.gen_sections,
                                st.session_state.gen_metadata,
                                st.session_state.gen_sources,
                            )
                        _gen_persist()
                        st.rerun()
                    except GeneratorAPIError as e:
                        st.error(f"Compliance review failed: {e}")
                    except Exception as e:
                        st.error(f"Compliance review failed unexpectedly: {e}")
        else:
            st.success("No issues flagged.")

        with st.expander("Suggested abbreviations", expanded=False):
            abbrevs = comp.get("abbreviations") or []
            if abbrevs:
                st.dataframe(pd.DataFrame(abbrevs), use_container_width=True, hide_index=True)
            else:
                st.info("None suggested.")

        with st.expander("Suggested footnotes", expanded=False):
            fns = comp.get("footnotes") or []
            if fns:
                for f in fns:
                    st.markdown(f"**{f.get('term', '')}** — {f.get('footnote_text', '')}")
            else:
                st.info("None suggested.")

        with st.expander("References", expanded=False):
            refs = comp.get("references") or []
            if refs:
                for r in refs:
                    st.write(f"- {r}")
            else:
                st.info("None compiled.")

        # Gate progression: compliance must have been run.
        compliance_run = bool(st.session_state.gen_compliance)
        if not compliance_run:
            st.warning("Run the compliance checks before continuing.")
        _gen_nav(prev_to=6, next_to=8 if compliance_run else None,
                 next_label="Continue to teaching note & export →")

    # ----- Step 8: teaching note & export --------------------------------
    elif st.session_state.gen_step == 8:
        st.header("Step 8 — Teaching note and export")
        ctype = st.session_state.gen_case_type
        meta = st.session_state.gen_metadata or {}

        if meta.get("teaching_note", True):
            st.subheader("Teaching note")
            if not st.session_state.gen_teaching_note or st.button("Regenerate teaching note", key="regen_tn"):
                with st.spinner("Generating teaching note (AGK template, 12.1-12.7)..."):
                    try:
                        st.session_state.gen_teaching_note = generate_teaching_note(
                            ctype,
                            st.session_state.gen_sections,
                            st.session_state.gen_metadata,
                            st.session_state.gen_sources,
                            st.session_state.gen_processed,
                        )
                        _gen_persist()
                    except GeneratorAPIError as e:
                        st.error(f"Teaching note generation failed: {e}")
                    except Exception as e:
                        st.error(f"Teaching note generation failed unexpectedly: {e}")
            tn = st.session_state.gen_teaching_note or {}
            if tn.get("case_summary"):
                st.markdown(f"**Case summary**\n\n{tn['case_summary']}")
            if tn.get("learning_objectives"):
                st.markdown("**Learning objectives**")
                for lo in tn["learning_objectives"]:
                    st.write(f"- {lo}")
            if tn.get("teaching_strategy"):
                with st.expander("Teaching strategy", expanded=False):
                    st.write(tn["teaching_strategy"])
            if tn.get("discussion_questions"):
                with st.expander("Discussion questions with answers", expanded=False):
                    for i, dq in enumerate(tn["discussion_questions"], 1):
                        st.markdown(f"**Q{i}.** {dq.get('question', '')}")
                        st.write(dq.get("answer_guide", ""))
                        st.markdown("---")
            if tn.get("assessment_methodology"):
                with st.expander("Assessment methodology", expanded=False):
                    st.write(tn["assessment_methodology"])
            if tn.get("target_audience_and_domains"):
                with st.expander("Target audience and subject domains", expanded=False):
                    st.write(tn["target_audience_and_domains"])
            if tn.get("additional_resources"):
                with st.expander("Additional teaching resources", expanded=False):
                    for r in tn["additional_resources"]:
                        st.markdown(f"- **{r.get('title', '')}** ({r.get('source', '')}) — {r.get('relevance', '')}")
        else:
            st.info("Teaching note generation was disabled in metadata.")

        st.markdown("---")
        st.subheader("Karmayogi Competency Model (KCM) mapping")
        st.caption(
            "Optional. Click the button below to weave KCM behavioural and "
            "functional competencies through the draft. Mapping does not run "
            "automatically."
        )
        kcm_btn_label = (
            "Re-map KCM competencies"
            if st.session_state.gen_kcm_mapping
            else "Map KCM competencies"
        )
        if st.button(kcm_btn_label, key="run_kcm"):
            with st.spinner("Mapping case to KCM behavioural and functional competencies…"):
                try:
                    st.session_state.gen_kcm_mapping = map_kcm_competencies(
                        ctype,
                        st.session_state.gen_sections,
                        st.session_state.gen_metadata,
                    )
                    _gen_persist()
                except GeneratorAPIError as e:
                    st.error(f"KCM mapping failed: {e}")
                except Exception as e:
                    st.error(f"KCM mapping failed unexpectedly: {e}")
        kcm = st.session_state.gen_kcm_mapping or {}
        comps = kcm.get("competencies", []) or []
        if comps:
            behavioural = [c for c in comps if str(c.get("category", "")).lower().startswith("behav")]
            functional = [c for c in comps if str(c.get("category", "")).lower().startswith("func")]
            kc_cols = st.columns(2)
            with kc_cols[0]:
                st.markdown("**Behavioural competencies**")
                if behavioural:
                    for c in behavioural:
                        st.markdown(f"- **{c.get('name', '')}** — {c.get('justification', '')}")
                else:
                    st.caption("None identified.")
            with kc_cols[1]:
                st.markdown("**Functional competencies**")
                if functional:
                    for c in functional:
                        st.markdown(f"- **{c.get('name', '')}** — {c.get('justification', '')}")
                else:
                    st.caption("None identified.")
            ws = kcm.get("weaving_suggestions", []) or []
            if ws:
                with st.expander("Suggestions for weaving these competencies more visibly", expanded=False):
                    for s in ws:
                        st.markdown(f"- {s}")

        st.markdown("---")
        st.subheader("Download")

        try:
            docx_bytes = build_case_docx(
                ctype,
                st.session_state.gen_metadata,
                st.session_state.gen_sections,
                compliance=st.session_state.gen_compliance,
            )
        except Exception as e:
            docx_bytes = None
            st.error(f"DOCX export failed: {e}")

        try:
            pdf_bytes = build_case_pdf(
                ctype,
                st.session_state.gen_metadata,
                st.session_state.gen_sections,
                compliance=st.session_state.gen_compliance,
            )
        except Exception as e:
            pdf_bytes = None
            st.error(f"PDF export failed: {e}")

        try:
            txt_bytes = build_case_txt(
                ctype,
                st.session_state.gen_metadata,
                st.session_state.gen_sections,
            )
            if isinstance(txt_bytes, str):
                txt_bytes = txt_bytes.encode("utf-8")
        except Exception as e:
            txt_bytes = None
            st.error(f"Plain-text export failed: {e}")

        tn_docx = None
        if st.session_state.gen_teaching_note:
            try:
                tn_docx = build_teaching_note_docx(
                    st.session_state.gen_metadata,
                    st.session_state.gen_teaching_note,
                    ctype,
                )
            except Exception as e:
                st.error(f"Teaching note export failed: {e}")

        title_for_file = re.sub(r"[^A-Za-z0-9_-]+", "_", (meta.get("title") or "case_study"))[:60] or "case_study"

        cols = st.columns(4)
        if docx_bytes:
            cols[0].download_button(
                "Case (DOCX)",
                data=docx_bytes,
                file_name=f"{title_for_file}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )
        if pdf_bytes:
            cols[1].download_button(
                "Case (PDF)",
                data=pdf_bytes,
                file_name=f"{title_for_file}.pdf",
                mime="application/pdf",
                use_container_width=True,
            )
        if txt_bytes:
            cols[2].download_button(
                "Case (TXT)",
                data=txt_bytes,
                file_name=f"{title_for_file}.txt",
                mime="text/plain",
                use_container_width=True,
            )
        if tn_docx:
            cols[3].download_button(
                "Teaching note (DOCX)",
                data=tn_docx,
                file_name=f"{title_for_file}_teaching_note.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )

        st.markdown("---")
        if st.button("Mark as finalised and save", key="gen_finalise"):
            try:
                # Persist final document bytes (base64-encoded) so they can be
                # re-downloaded later without regenerating from the AI calls.
                final_docs = {}
                if docx_bytes:
                    final_docs["docx_b64"] = base64.b64encode(docx_bytes).decode("ascii")
                if pdf_bytes:
                    final_docs["pdf_b64"] = base64.b64encode(pdf_bytes).decode("ascii")
                if txt_bytes:
                    final_docs["txt_b64"] = base64.b64encode(txt_bytes).decode("ascii")
                if tn_docx:
                    final_docs["teaching_note_docx_b64"] = base64.b64encode(tn_docx).decode("ascii")
                final_docs["finalised_at"] = datetime.utcnow().isoformat() + "Z"
                final_docs["filename_stem"] = title_for_file
                st.session_state.gen_final_documents = final_docs

                cid = save_generated_case(
                    user_id=st.session_state.user_id,
                    case_id=st.session_state.gen_case_id,
                    title=meta.get("title") or "Untitled",
                    case_type=ctype,
                    status="finalised",
                    metadata=st.session_state.gen_metadata,
                    intent=st.session_state.gen_intent,
                    sources=st.session_state.gen_sources,
                    processed=st.session_state.gen_processed,
                    sections=st.session_state.gen_sections,
                    compliance=st.session_state.gen_compliance,
                    teaching_note=st.session_state.gen_teaching_note,
                    kcm_mapping=st.session_state.gen_kcm_mapping,
                    final_documents=final_docs,
                )
                if cid:
                    st.session_state.gen_case_id = cid
                    st.success("Saved as finalised. Final DOCX/PDF/TXT are stored with the draft.")
                else:
                    st.error("Could not save.")
            except Exception as e:
                st.error(f"Save failed: {e}")

        _gen_nav(prev_to=7, next_to=None)

elif st.session_state.get("active_tool") == "analyser" and st.session_state.case_study_text:
    # Document info section
    st.header("Document Information")
    st.write(f"**Filename:** {st.session_state.document_name}")
    
    # Summary section
    st.header("Case Study Summary")
    if st.session_state.case_study_summary:
        st.write(st.session_state.case_study_summary)
    else:
        st.info("Click 'Generate Summary' to create a summary of the case study document")
    
    if st.session_state.get('sector_tags'):
        st.header("Sector Tags & Keywords")
        
        sector_badges = " ".join([
            f'<span style="display: inline-block; padding: 5px 14px; margin: 4px; background-color: #DBEAFE; color: #1E40AF; border-radius: 20px; font-weight: 600; font-size: 0.9em;">{tag}</span>'
            for tag in st.session_state.sector_tags
        ])
        st.markdown(f"**Sector Tags:** {sector_badges}", unsafe_allow_html=True)
        
        subthemes = st.session_state.get('sector_subthemes', {})
        if subthemes:
            st.markdown("**Sub-themes:**")
            for sector, themes in subthemes.items():
                if themes:
                    theme_badges = " ".join([
                        f'<span style="display: inline-block; padding: 3px 10px; margin: 3px; background-color: #E0F2FE; color: #0369A1; border-radius: 15px; font-size: 0.85em;">{t}</span>'
                        for t in themes
                    ])
                    st.markdown(f"&nbsp;&nbsp;*{sector}:* {theme_badges}", unsafe_allow_html=True)
        
        keywords = st.session_state.get('keywords', [])
        if keywords:
            keyword_badges = " ".join([
                f'<span style="display: inline-block; padding: 4px 12px; margin: 3px; background-color: #F3F4F6; color: #374151; border-radius: 15px; font-size: 0.85em;">{kw}</span>'
                for kw in keywords
            ])
            st.markdown(f"**Keywords:** {keyword_badges}", unsafe_allow_html=True)
        
        st.markdown("<br>", unsafe_allow_html=True)
    
    # Assessment Results
    if st.session_state.assessment_results:
        st.header("Assessment Results")
        
        # Display Overall Composite Score prominently
        if st.session_state.weighted_scores:
            ws = st.session_state.weighted_scores
            final_score = ws["final_composite_score"]
            grade_label = get_grade_label(final_score)
            
            # Create a prominent score display
            col1, col2, col3 = st.columns([0.1, 0.8, 0.1])
            with col2:
                # Determine color based on grade
                if grade_label == "Excellent":
                    color = "#28a745"
                    badge_class = "score-excellent"
                elif grade_label == "Good":
                    color = "#3B82F6"
                    badge_class = "score-good"
                elif grade_label == "Satisfactory":
                    color = "#ffc107"
                    badge_class = "score-satisfactory"
                elif grade_label == "Needs Improvement":
                    color = "#fd7e14"
                    badge_class = "score-needs-improvement"
                else:
                    color = "#dc3545"
                    badge_class = "score-poor"
                
                st.markdown(f"""
                <div style="text-align: center; padding: 20px; background: linear-gradient(135deg, #f5f7fa 0%, #c3cfe2 100%); border-radius: 15px; margin-bottom: 20px;">
                    <h2 style="margin: 0; color: #1E3A8A;">Final Composite Score</h2>
                    <h1 style="font-size: 3.5em; margin: 10px 0; color: {color};">{final_score:.1f}%</h1>
                    <span class="score-badge {badge_class}" style="font-size: 1.2em; padding: 8px 20px;">{grade_label}</span>
                </div>
                """, unsafe_allow_html=True)
            
            # Display weighted area breakdown
            st.subheader("Weighted Area Scores")
            
            # Create columns for area scores
            area_cols = st.columns(4)
            for i, (area_id, area_info) in enumerate(ASSESSMENT_AREAS.items()):
                with area_cols[i]:
                    area_score_data = ws["area_scores"].get(area_id, {})
                    weighted_data = ws["weighted_scores"].get(area_id, {})
                    
                    score = area_score_data.get("score", 0)
                    max_score = area_score_data.get("max_score", 1)
                    percentage = area_score_data.get("percentage", 0)
                    weight = weighted_data.get("weight", 0) * 100
                    contribution = weighted_data.get("weighted_contribution", 0) * 100
                    
                    st.markdown(f"""
                    <div style="text-align: center; padding: 15px; background-color: #F9FAFB; border-radius: 10px; border: 1px solid #E5E7EB; min-height: 250px; display: flex; flex-direction: column; justify-content: space-between;">
                        <div>
                            <p style="font-size: 0.85em; color: #6B7280; margin-bottom: 10px; line-height: 1.2; font-weight: 500;">{area_info['name']}</p>
                            <h3 style="margin: 10px 0; color: #1E3A8A; font-size: 1.6em;">{score}/{max_score}</h3>
                            <p style="font-size: 0.95em; color: #4B5563; margin: 0; font-weight: 500;">({percentage:.1f}%)</p>
                        </div>
                        <div style="border-top: 1px solid #E5E7EB; padding-top: 10px; margin-top: 10px;">
                            <p style="font-size: 0.8em; color: #9CA3AF; margin-bottom: 2px;">Weight: {weight:.0f}%</p>
                            <p style="font-size: 0.85em; color: #059669; margin: 0; font-weight: 600;">Contribution: {contribution:.1f}%</p>
                        </div>
                    </div>
                    """, unsafe_allow_html=True)
        
        # Action buttons section
        st.markdown("<br>", unsafe_allow_html=True)
        col1, col2 = st.columns(2)
        
        with col1:
            # Generate recommendations for case study improvement
            if 'recommendations' not in st.session_state or not st.session_state.recommendations:
                with st.spinner("Generating recommendations..."):
                    # Prepare data for recommendations generation
                    low_scores = []
                    for area_id, area_results in st.session_state.assessment_results.items():
                        area_name = ASSESSMENT_AREAS[area_id]["name"]
                        for criterion_id, result in area_results.items():
                            criterion_info = ASSESSMENT_CRITERIA[area_id][criterion_id]
                            if criterion_info.get("informational", False):
                                continue
                            criterion_name = criterion_info["name"]
                            max_score = criterion_info.get("max_score", 3)
                            score = result.get("score", 0)
                            
                            if score < max_score * 0.5:
                                low_scores.append({
                                    "area": area_name,
                                    "criterion": criterion_name,
                                    "score": score,
                                    "max_score": max_score,
                                    "reasoning": result.get("reasoning", "")
                                })
                    
                    # Generate recommendations based on low scores
                    if low_scores:
                        prompt = f"""
                        Based on the case study assessment using the CBC-India AGK Case Study Review Rubric, 
                        please generate 5-8 concrete, actionable recommendations to improve the case study.
                        
                        Focus on addressing the following areas with low scores:
                        
                        {json.dumps(low_scores, indent=2)}
                        
                        Format each recommendation as a numbered entry starting with "Recommendation 1:", 
                        "Recommendation 2:", etc. Each recommendation should be a complete, self-contained 
                        paragraph without line breaks or bullets in the middle of sentences.
                        
                        {PROMPT_EXCLUSION_INSTRUCTIONS}
                        
                        Your recommendations should be:
                        1. Specific and actionable for case study improvement
                        2. Practical to implement
                        3. Directly address deficiencies identified in the assessment
                        4. Aligned with the AGK case study writing standards
                        5. Written in complete sentences with proper punctuation
                        
                        Example format:
                        Recommendation 1: [Complete sentence describing specific action] to address [specific issue]. 
                        The recommendation should include enough detail to be actionable but be contained in a single complete paragraph.
                        """
                        
                        st.session_state.recommendations = call_openai_api(prompt)
                    else:
                        st.session_state.recommendations = "The case study generally scores well across all assessment areas. Consider maintaining the current approach while reviewing any minor areas for potential enhancement."
            
            # Save to history button (only shown to logged-in users)
            if st.session_state.logged_in:
                if st.button("💾 Save to History", use_container_width=True):
                    if st.session_state.user_id:
                        # Save assessment to database
                        try:
                            results_to_save = dict(st.session_state.assessment_results)
                            results_to_save['_sector_tags'] = st.session_state.get('sector_tags', [])
                            results_to_save['_sector_subthemes'] = st.session_state.get('sector_subthemes', {})
                            results_to_save['_keywords'] = st.session_state.get('keywords', [])
                            assessment_id = save_assessment(
                                user_id=st.session_state.user_id,
                                document_name=st.session_state.document_name,
                                policy_summary=st.session_state.case_study_summary,
                                assessment_results=results_to_save,
                                recommendations=st.session_state.get('recommendations', "No recommendations available.")
                            )
                            
                            if assessment_id:
                                # Make success message more prominent
                                st.success(f"Assessment successfully saved to your history!")
                            else:
                                st.error("Failed to save assessment. Database error occurred.")
                        except Exception as e:
                            st.error(f"Error saving assessment: {str(e)}")
                            st.info("Please try again or check database connection.")
                    else:
                        st.error("Error saving assessment. Please try again.")
        
        with col2:
            # Export as PDF button
            st.download_button(
                label="📄 Export Report as PDF",
                data=generate_report_pdf(
                    "case_study_assessment.pdf",
                    st.session_state.document_name,
                    st.session_state.case_study_summary,
                    st.session_state.assessment_results,
                    ASSESSMENT_AREAS,
                    ASSESSMENT_CRITERIA,
                    st.session_state.get('recommendations', "No recommendations available."),
                    st.session_state.weighted_scores,
                    st.session_state.get('competency_mapping', None),
                    tags_data={"sectors": st.session_state.get('sector_tags', []), "subthemes": st.session_state.get('sector_subthemes', {}), "keywords": st.session_state.get('keywords', [])},
                    writing_findings=st.session_state.get('writing_findings', [])
                ),
                file_name="case_study_assessment.pdf",
                mime="application/pdf",
                use_container_width=True
            )
        
        st.markdown("<br>", unsafe_allow_html=True)
        
        # Create tabs for each assessment area
        tabs = st.tabs([ASSESSMENT_AREAS[area_id]["name"] for area_id in ASSESSMENT_AREAS])
        
        for i, (area_id, area_info) in enumerate(ASSESSMENT_AREAS.items()):
            with tabs[i]:
                st.subheader(area_info["name"])
                st.write(area_info["description"])
                st.write(f"**Weight:** {area_info['weight'] * 100:.0f}% | **Total Points:** {area_info['total_points']}")
                
                if area_id in st.session_state.assessment_results:
                    area_criteria = ASSESSMENT_CRITERIA[area_id]
                    area_results = st.session_state.assessment_results[area_id]
                    
                    total_score = sum(
                        result.get("score", 0) for crit_id, result in area_results.items()
                        if not ASSESSMENT_CRITERIA.get(area_id, {}).get(crit_id, {}).get("informational", False)
                    )
                    max_possible = area_info["total_points"]
                    percentage = (total_score / max_possible * 100) if max_possible > 0 else 0
                    
                    # Display scores in a metrics row
                    col1, col2, col3 = st.columns(3)
                    with col1:
                        st.metric("Area Score", f"{total_score}/{max_possible}")
                    with col2:
                        st.metric("Percentage", f"{percentage:.1f}%")
                    with col3:
                        weighted_contribution = (area_info["weight"] * (total_score / max_possible) * 100) if max_possible > 0 else 0
                        st.metric("Weighted Contribution", f"{weighted_contribution:.1f}%")
                    
                    # Create visualizations for scores
                    st.subheader("Score Visualization")
                    
                    # Create gauge chart for area percentage
                    gauge_fig = create_gauge_chart(percentage, f"Score for {area_info['name']}", max_value=100)
                    st.plotly_chart(gauge_fig)
                    
                    criteria_names = []
                    criteria_scores = []
                    criteria_max_scores = []
                    
                    for crit_id, result in area_results.items():
                        if area_criteria[crit_id].get("informational", False):
                            continue
                        criteria_names.append(area_criteria[crit_id]["name"])
                        criteria_scores.append(result.get("score", 0))
                        criteria_max_scores.append(area_criteria[crit_id].get("max_score", 3))
                    
                    # Create DataFrame for the chart
                    df = pd.DataFrame({
                        'Criteria': criteria_names,
                        'Score': criteria_scores,
                        'Max Score': criteria_max_scores
                    })
                    
                    fig = go.Figure()
                    
                    # Add bars for actual scores
                    fig.add_trace(go.Bar(
                        x=df['Criteria'],
                        y=df['Score'],
                        name='Score',
                        marker_color=['#28a745' if s/m >= 0.75 else '#ffc107' if s/m >= 0.5 else '#dc3545' 
                                      for s, m in zip(df['Score'], df['Max Score'])],
                        text=[f"{s}/{m}" for s, m in zip(df['Score'], df['Max Score'])],
                        textposition='outside'
                    ))
                    
                    fig.update_layout(
                        title=f"Criteria Scores for {area_info['name']}",
                        xaxis_title="Criteria",
                        yaxis_title="Score",
                        height=400,
                        showlegend=False
                    )
                    
                    st.plotly_chart(fig)
                    
                    # Detailed criteria analysis
                    st.subheader("Detailed Assessment")
                    
                    for criterion_id, result in area_results.items():
                        criterion_info = area_criteria[criterion_id]
                        criterion_name = criterion_info["name"]
                        is_info = criterion_info.get("informational", False)
                        doc_ref = result.get("document_reference", "N/A")
                        
                        if is_info:
                            narrative = result.get("narrative", result.get("reasoning", "N/A"))
                            with st.container():
                                st.markdown(f"""
                                <div style="background-color: #EFF6FF; padding: 15px; border-radius: 10px; margin-bottom: 15px; border-left: 4px solid #3B82F6;">
                                    <h4 style="margin: 0 0 10px 0; color: #1E3A8A;">{criterion_name}</h4>
                                    <div style="display: inline-block; padding: 5px 15px; background-color: #3B82F6; color: white; border-radius: 20px; font-weight: bold; font-size: 0.85em;">
                                        ℹ Informational
                                    </div>
                                </div>
                                """, unsafe_allow_html=True)
                                
                                st.write(f"**Analysis:** {narrative}")
                                st.write(f"**Document Reference:** {doc_ref}")
                                st.markdown("---")
                        else:
                            max_score = criterion_info.get("max_score", 3)
                            scoring_logic = criterion_info.get("scoring_logic", "")
                            score = result.get("score", 0)
                            reasoning = result.get("reasoning", "N/A")
                            
                            score_pct = score / max_score if max_score > 0 else 0
                            if score_pct >= 0.75:
                                score_color = "#28a745"
                            elif score_pct >= 0.5:
                                score_color = "#ffc107"
                            else:
                                score_color = "#dc3545"
                            
                            with st.container():
                                st.markdown(f"""
                                <div style="background-color: #F9FAFB; padding: 15px; border-radius: 10px; margin-bottom: 15px; border-left: 4px solid {score_color};">
                                    <h4 style="margin: 0 0 10px 0; color: #1E3A8A;">{criterion_name}</h4>
                                    <p style="font-size: 0.9em; color: #6B7280; margin-bottom: 10px;"><strong>Scoring:</strong> {scoring_logic}</p>
                                    <div style="display: inline-block; padding: 5px 15px; background-color: {score_color}; color: white; border-radius: 20px; font-weight: bold;">
                                        Score: {score}/{max_score}
                                    </div>
                                </div>
                                """, unsafe_allow_html=True)
                                
                                st.write(f"**Reasoning & Evidence:** {reasoning}")
                                st.write(f"**Document Reference:** {doc_ref}")
                                st.markdown("---")

                    if area_id == "area2":
                        writing_findings = st.session_state.get('writing_findings', [])
                        if writing_findings:
                            st.subheader("Writing Assistant Findings")
                            st.write("Detailed analysis of grammar, spelling, tense consistency, redundancy, and sentence structure across the complete case study.")

                            type_counts = {}
                            severity_counts = {"High": 0, "Medium": 0, "Low": 0}
                            for f in writing_findings:
                                t = f.get("type", "Other")
                                s = f.get("severity", "Low")
                                type_counts[t] = type_counts.get(t, 0) + 1
                                severity_counts[s] = severity_counts.get(s, 0) + 1

                            summary_cols = st.columns(len(type_counts) + 1)
                            with summary_cols[0]:
                                st.metric("Total Issues", len(writing_findings))
                            for idx, (issue_type, count) in enumerate(type_counts.items()):
                                with summary_cols[idx + 1]:
                                    st.metric(issue_type, count)

                            sev_cols = st.columns(3)
                            sev_colors = {"High": "#dc3545", "Medium": "#ffc107", "Low": "#6c757d"}
                            for idx, (sev, count) in enumerate(severity_counts.items()):
                                with sev_cols[idx]:
                                    st.markdown(f"""
                                    <div style="text-align:center; padding:8px; background-color:{sev_colors[sev]}20; border-radius:8px; border:1px solid {sev_colors[sev]};">
                                        <span style="color:{sev_colors[sev]}; font-weight:bold;">{sev}: {count}</span>
                                    </div>
                                    """, unsafe_allow_html=True)

                            st.markdown("<br>", unsafe_allow_html=True)

                            table_data = []
                            for i, f in enumerate(writing_findings, 1):
                                table_data.append({
                                    "Issue #": i,
                                    "Type": f.get("type", "Other"),
                                    "Original Text": f.get("original_text", ""),
                                    "Suggested Fix": f.get("suggestion", ""),
                                    "Severity": f.get("severity", "Low"),
                                    "Explanation": f.get("context", "")
                                })

                            findings_df = pd.DataFrame(table_data)

                            def style_severity(val):
                                colors = {"High": "background-color: #f8d7da; color: #721c24;", "Medium": "background-color: #fff3cd; color: #856404;", "Low": "background-color: #e2e3e5; color: #383d41;"}
                                return colors.get(val, "")

                            styled_df = findings_df.style.applymap(style_severity, subset=["Severity"])
                            st.dataframe(styled_df, use_container_width=True, hide_index=True, height=min(400, 50 + len(table_data) * 35))

                            with st.expander("View All Findings in Detail", expanded=False):
                                for i, f in enumerate(writing_findings, 1):
                                    sev = f.get("severity", "Low")
                                    sev_color = sev_colors.get(sev, "#6c757d")
                                    type_badge_colors = {"Spelling": "#0d6efd", "Grammar": "#6610f2", "Tense": "#fd7e14", "Redundancy": "#20c997", "Structure": "#0dcaf0"}
                                    type_color = type_badge_colors.get(f.get("type", ""), "#6c757d")
                                    st.markdown(f"""
                                    <div style="padding:12px; margin-bottom:8px; background-color:#f8f9fa; border-radius:8px; border-left:4px solid {sev_color};">
                                        <div style="margin-bottom:6px;">
                                            <span style="display:inline-block; padding:2px 10px; background-color:{type_color}; color:white; border-radius:12px; font-size:0.8em; font-weight:bold;">{f.get("type", "Other")}</span>
                                            <span style="display:inline-block; padding:2px 10px; background-color:{sev_color}20; color:{sev_color}; border-radius:12px; font-size:0.8em; font-weight:bold; margin-left:5px;">{sev}</span>
                                        </div>
                                        <p style="margin:4px 0;"><strong>Original:</strong> <span style="color:#dc3545; text-decoration:line-through;">{f.get("original_text", "")}</span></p>
                                        <p style="margin:4px 0;"><strong>Suggested:</strong> <span style="color:#28a745;">{f.get("suggestion", "")}</span></p>
                                        <p style="margin:4px 0; font-size:0.9em; color:#6B7280;"><em>{f.get("context", "")}</em></p>
                                    </div>
                                    """, unsafe_allow_html=True)

                    # Display KCM Competency Mapping before recommendations
                    competency_data = st.session_state.get('competency_mapping')
                    if competency_data and isinstance(competency_data, dict):
                        st.subheader("Karmayogi Competency Mapping")
                        st.write("Based on the case study analysis, the following competencies from the Karmayogi Competency Model (KCM) are most relevant:")
                        
                        # Behavioral Competencies
                        behavioral_comps = competency_data.get('behavioral_competencies', [])
                        if behavioral_comps and isinstance(behavioral_comps, list):
                            st.markdown("#### Behavioral Sub-competencies")
                            for comp in behavioral_comps:
                                if isinstance(comp, dict):
                                    comp_name = comp.get('name', 'Unknown')
                                    parent = comp.get('parent_competency', '')
                                    display_name = f"{comp_name} ({parent})" if parent else comp_name
                                    justification = comp.get('justification', 'No justification provided')
                                    st.markdown(f"""
                                    <div style="background-color: #E0F2FE; padding: 12px; border-radius: 8px; margin-bottom: 10px; border-left: 4px solid #0284C7;">
                                        <strong style="color: #0369A1;">{display_name}</strong>
                                        <p style="margin: 8px 0 0 0; color: #374151; font-size: 0.95em;">{justification}</p>
                                    </div>
                                    """, unsafe_allow_html=True)
                        
                        # Functional Competencies
                        functional_comps = competency_data.get('functional_competencies', [])
                        if functional_comps and isinstance(functional_comps, list):
                            st.markdown("#### Functional Sub-competencies")
                            for comp in functional_comps:
                                if isinstance(comp, dict):
                                    comp_name = comp.get('name', 'Unknown')
                                    parent = comp.get('parent_competency', '')
                                    display_name = f"{comp_name} ({parent})" if parent else comp_name
                                    justification = comp.get('justification', 'No justification provided')
                                    st.markdown(f"""
                                    <div style="background-color: #ECFDF5; padding: 12px; border-radius: 8px; margin-bottom: 10px; border-left: 4px solid #059669;">
                                        <strong style="color: #047857;">{display_name}</strong>
                                        <p style="margin: 8px 0 0 0; color: #374151; font-size: 0.95em;">{justification}</p>
                                    </div>
                                    """, unsafe_allow_html=True)
                        
                        st.markdown("---")
                    
                    # After all criteria are shown, add recommendations section
                    if 'recommendations' in st.session_state:
                        st.subheader("Recommendations for Improvement")
                        st.write("Based on the assessment, we recommend the following:")
                        
                        # Format recommendations as a simple list
                        recommendations = st.session_state.recommendations
                        
                        # Process the recommendations text
                        if "Recommendation" in recommendations:
                            # Split by "Recommendation" and process each item
                            if recommendations.strip().startswith("Recommendation"):
                                parts = recommendations.split("Recommendation")
                                for part in parts:
                                    if part.strip():
                                        # Clean up the format
                                        cleaned_part = part.strip().replace('\n', ' ')
                                        
                                        # Look for colons to preserve formatting
                                        if ':' in cleaned_part and cleaned_part.find(':') < 10:
                                            num, content = cleaned_part.split(':', 1)
                                            st.write(f"**Recommendation {num.strip()}:** {content.strip()}")
                                        else:
                                            st.write(f"**Recommendation {cleaned_part.strip()}**")
                        
                        # Fallback to regular dash-based splitting if no "Recommendation" keyword
                        elif "-" in recommendations:
                            rec_points = recommendations.split("-")
                            for point in rec_points:
                                if point.strip():
                                    st.write(f"- {point.strip()}")
                        else:
                            # Just show as regular text if no structured format detected
                            st.write(recommendations)
                else:
                    st.info("No assessment data available for this area")
    
    # Raw extracted text (collapsible)
    with st.expander("View Extracted Text"):
        st.subheader("Case Study Document")
        st.text_area("Case Study Text", st.session_state.case_study_text, height=250, key="view_case_study")
        if st.session_state.teaching_note_text:
            st.subheader("Teaching Note Document")
            st.text_area("Teaching Note Text", st.session_state.teaching_note_text, height=250, key="view_teaching_note")
elif st.session_state.get("active_tool") == "analyser" and not st.session_state.case_study_text:
    st.info("Please upload both Case Study and Teaching Note documents using the sidebar to begin analysis.")

elif not st.session_state.logged_in:
    st.info("Please login to access the tools.")
